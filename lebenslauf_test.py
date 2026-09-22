# -*- coding: utf-8 -*-
"""Selbsttest des Lebenslauf-Bauers – läuft gegen eine Kopie der Datenbank.

    python -X utf8 lebenslauf_test.py

Geprüft wird der Weg, den die Kollegin im Alltag geht: Kunde in der Akte öffnen,
„Lebenslauf bauen", Formular ausfüllen, Excel herunterladen. Braucht kein Internet.
"""
import os
import re
import sys

HIER = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HIER)
import pruefkopie                # noqa: E402
# Warum die Arbeitskopie über `sqlite3.backup` läuft und nicht über `shutil.copy`,
# steht im Kopf von `pruefkopie.py`. Am Echtbestand ändert der Lauf nichts.
kopie = pruefkopie.anlegen("improfy_os_cv_test.db")
os.environ["IMPROFY_OS_DB"] = kopie

import openpyxl                     # noqa: E402
from flask import render_template as flask_render   # noqa: E402
import app as A                     # noqa: E402
import datenbank as db              # noqa: E402
import lebenslauf_bauen as LB       # noqa: E402

ergebnis = []


def vorlagen_dateien():
    ordner = os.path.join(HIER, "templates", "cv_designs")
    return [os.path.join(ordner, f) for f in os.listdir(ordner)
            if f.startswith("cv_") and f.endswith(".html")]


def offen(pfad):
    with open(pfad, encoding="utf-8") as f:
        return f.read()


def ohne_kommentare(text):
    return re.sub(r"/\*.*?\*/", "", text, flags=re.S)


def erster_absender(html):
    """Der erste Absende-Knopf im grossen Formular – der, den Enter drückt.

    HTML macht den ersten Absende-Knopf eines Formulars zum Standardknopf: Ein Enter in
    irgendeinem einzeiligen Feld löst ihn aus. Welcher das ist, entscheidet allein die
    Reihenfolge im Blatt – deshalb wird hier das HTML gelesen und nicht geklickt."""
    anfang = html.find('<form method="post" enctype="multipart/form-data">')
    if anfang < 0:
        return ""
    innen = html[anfang:html.find("</form>", anfang)]
    for treffer in re.finditer(r"<(?:button|input)\b[^>]*>", innen):
        if 'type="submit"' in treffer.group(0):
            return treffer.group(0)
    return ""


def _fotos_je_seite(rohdaten):
    """Wie viele Bewerbungsfotos auf jeder Seite des PDF stehen – eine Zahl je Seite.

    Gezaehlt wird mit derselben Erkennung wie in `seitenpruefung` (hochkante Bildbloecke),
    damit Pruefung und Selbsttest nicht zwei verschiedene Dinge „Foto" nennen."""
    import pymupdf
    import seitenpruefung as SP
    with pymupdf.open(stream=rohdaten, filetype="pdf") as doc:
        return [len(SP._fotos(s)) for s in doc]


def _fotobreite_mm(rohdaten):
    """Breite des Fotos auf Seite 1 in Millimetern, 0.0 wenn keines da ist.

    Das ist das Mass, an dem eine Stauchung auffaellt: Chrome verkleinert bei einem
    Element ueber der Blattkante **das ganze Dokument**, das Blatt bleibt aber A4."""
    import pymupdf
    import seitenpruefung as SP
    with pymupdf.open(stream=rohdaten, filetype="pdf") as doc:
        fotos = SP._fotos(doc[0])
        return (fotos[0][2] - fotos[0][0]) / SP.MM if fotos else 0.0


def pruefe(name, bedingung, detail=""):
    ergebnis.append((name, bool(bedingung)))
    print(f"  {'OK  ' if bedingung else 'FEHL'} {name}{(' – ' + str(detail)) if detail else ''}")


def main():
    db.init()
    c = A.app.test_client()
    kid = db.wert("SELECT kunde_id FROM gutschein_zeile WHERE kunde_id IS NOT NULL"
                  "   AND improfy_id IS NOT NULL AND improfy_id<>'' LIMIT 1")
    kunde = db.eine("SELECT * FROM kunde WHERE id=?", (kid,))
    print(f"\nTestkunde: {kunde['name']} (id {kid})")

    print("\n1. Vorbelegung aus dem OS")
    v = LB.vorbelegung(kid)
    pruefe("Kunde gefunden, Name geteilt", v and v["daten"]["vorname"] and v["daten"]["nachname"],
           f"{v['daten']['vorname']} | {v['daten']['nachname']}")
    pruefe("Interne Kundennummer kommt aus dem Register, nicht aus der Tastatur",
           re.match(r"^ID-K\d+$", v["interne_id"] or ""), v["interne_id"])
    pruefe("Blattname ist ID + Name", v["blattname"].startswith(v["interne_id"]), v["blattname"])
    pruefe("Telefon aus der Kundenakte übernommen",
           (v["daten"]["mobil"] or "") == (kunde["telefon"] or ""), v["daten"]["mobil"])
    pruefe("Deutsch steht bei den Sprachen zuerst",
           v["daten"]["sprachen"] and v["daten"]["sprachen"][0]["sprache"] == "DEUTSCH")
    pruefe("Coach steht im Feld „Kunde von\"", bool(v["daten"]["kunde_von"]), v["daten"]["kunde_von"])

    print("\n2. Formularseite")
    r = c.get(f"/kunde/{kid}/lebenslauf")
    pruefe("Seite antwortet und zeigt den Blattnamen", r.status_code == 200 and v["blattname"] in r.text)
    pruefe("Vorbelegte Werte stehen im Formular", v["daten"]["nachname"] in r.text)

    print("\n3. Excel erzeugen")
    daten = {
        "vorname": v["daten"]["vorname"], "nachname": v["daten"]["nachname"],
        "angestrebter_job": "Lagerhelfer", "geschlecht": "m", "geburtsdatum": "01.01.1990",
        "mobil": "+49 170 0000000", "email": "test@example.de", "adresse": "Teststr. 1, 50667 Köln",
        "kunde_von": "Selbsttest", "fs_vorhanden": "on", "fs_klasse": "B", "fs_eu": "on",
        "beruf_zeitraum": ["01.2020 - 12.2023"], "beruf_firma": ["Testfirma GmbH, Köln"],
        "beruf_jobtitel": ["Lagerhelfer"], "beruf_taet": ["Kommissionieren\nWareneingang"],
        "bild_zeitraum": ["2015 - 2018"], "bild_abschluss": ["Berufsausbildung"],
        "bild_institution": ["Berufsschule"], "bild_note": [""],
        "spr_sprache": ["DEUTSCH", "ARABISCH"], "spr_niveau": ["gute Kenntnisse", "Muttersprache"],
        "edv_programm": ["MS Word"], "edv_sterne": ["4"],
        "ss_eigenschaft": ["Zuverlässigkeit"], "ss_sterne": ["5"],
        "zusatzqual": "Sachkunde §34a", "hobbys": "Fußball", "ueber_mich": "Guten Tag, ...",
    }
    r = c.post(f"/kunde/{kid}/lebenslauf", data=daten)
    pruefe("Excel kommt als Download zurück", r.status_code == 200 and len(r.data) > 20000,
           f"{len(r.data)} Bytes")
    dateiname = re.search(r'filename="([^"]+)"', r.headers.get("Content-Disposition", ""))
    dateiname = dateiname.group(1) if dateiname else ""
    pruefe("Dateiname trägt die Kunden-ID", dateiname.startswith(v["interne_id"]), dateiname)

    print("\n4. Inhalt der Datei")
    pfad = LB.datei(dateiname)
    ws = openpyxl.load_workbook(pfad).active
    pruefe("Blatt heißt wie der Kunde", ws.title.startswith(v["interne_id"]), ws.title)
    pruefe("Kopfdaten stehen in der Vorlage",
           ws["B10"].value == daten["vorname"] and ws["B12"].value == "Lagerhelfer"
           and ws["B14"].value == daten["mobil"], [ws[z].value for z in ("B10", "B12", "B14")])
    pruefe("Improfy-Block geschlechtsrichtig", ws["D25"].value == "Teilnehmer", ws["D25"].value)
    pruefe("Kunde von gesetzt", ws["B3"].value == "Selbsttest", ws["B3"].value)
    pruefe("Berufserfahrung eingetragen",
           ws["B30"].value == "01.2020 - 12.2023" and ws["D31"].value == "Lagerhelfer")
    pruefe("Bildung eingetragen", ws["D98"].value == "Berufsausbildung")

    print("\n5. Ergebnis hängt am Kunden")
    gebaut = LB.gebaute(kid)
    pruefe("Lebenslauf ist am Kunden vermerkt", any(g["name"] == dateiname for g in gebaut))
    dok = db.eine("SELECT vorhanden, dateiname FROM dokument WHERE kunde_id=? AND schluessel='lebenslauf'",
                  (kid,))
    pruefe("Qualitätsmanagement zählt den Lebenslauf als vorhanden",
           dok and dok["vorhanden"] == 1, dict(dok) if dok else None)
    pruefe("Datei lässt sich erneut herunterladen",
           c.get("/lebenslauf/datei/" + dateiname).status_code == 200)
    pruefe("Kundenakte verweist auf den Bauer",
           "lebenslauf" in c.get(f"/kunde/{kid}").text.casefold())

    print("\n6. Alte Unterlagen ablegen")
    import io
    import dokument_lesen as DL

    def feld(text, name):
        return re.findall('name="%s" value="([^"]*)"' % name, text)

    r = c.post(f"/kunde/{kid}/lebenslauf/lesen",
               data={"unterlagen": (io.BytesIO(open(pfad, "rb").read()), "alter_bogen.xlsx")},
               content_type="multipart/form-data")
    t = r.get_data(as_text=True)
    pruefe("Hochgeladene Improfy-Excel wird zurückgelesen",
           r.status_code == 200 and "Improfy-Vorlage gelesen" in t)
    pruefe("Alle Felder kommen aus der Excel zurück",
           feld(t, "angestrebter_job")[:1] == ["Lagerhelfer"]
           and feld(t, "beruf_firma")[:1] == ["Testfirma GmbH, Köln"]
           and feld(t, "bild_abschluss")[:1] == ["Berufsausbildung"]
           and feld(t, "fs_klasse")[:1] == ["B"],
           [feld(t, n)[:1] for n in ("angestrebter_job", "beruf_firma", "bild_abschluss", "fs_klasse")])
    pruefe("Tätigkeiten stehen wieder im Textfeld",
           "Kommissionieren" in t)

    text = ("Max Mustermann\n01.01.1990\nBERUFSERFAHRUNG\n"
            "(2020 - 2023) Lagerhelfer | Amazon GmbH,\nKöln\nWarenannahme\n"
            "SPRACHKENNTNISSE\nDEUTSCH - B1\n")
    daten2, gefunden2, hinweise2, roh2 = DL.lese_unterlagen(
        [("lebenslauf.txt", text.encode("utf-8"))])
    pruefe("Textdatei wird gelesen", bool(gefunden2), gefunden2)
    pruefe("Firma und Titel stehen richtig herum",
           daten2["berufserfahrung"][0]["firma"].startswith("Amazon")
           and daten2["berufserfahrung"][0]["jobtitel"] == "Lagerhelfer",
           daten2["berufserfahrung"][:1])
    pruefe("Umgebrochener Firmenname wird zusammengesetzt",
           daten2["berufserfahrung"][0]["firma"] == "Amazon GmbH, Köln",
           daten2["berufserfahrung"][0]["firma"])
    pruefe("Jedes Formularfeld ist im Ergebnis vorhanden",
           set(DL.leere_daten()) <= set(daten2))

    _, _, hinweise3, _ = DL.lese_unterlagen([("foto.jpg", b"\xff\xd8\xff")])
    pruefe("Bilder werden klar abgelehnt",
           any("Bilder" in h for h in hinweise3), hinweise3)
    _, _, hinweise4, _ = DL.lese_unterlagen([("scan.pdf", b"%PDF-1.4 kaputt")])
    pruefe("Unlesbares PDF wird gemeldet", bool(hinweise4), hinweise4)
    r = c.post(f"/kunde/{kid}/lebenslauf/lesen", data={"rohtext": ""})
    pruefe("Ohne Datei und ohne Text kommt eine Meldung, kein Absturz",
           r.status_code == 200 and "ließ sich nichts auslesen" in r.get_data(as_text=True))

    print("\n7. Designtes PDF")
    import cv_pdf
    block = cv_pdf.improfy_block({"geschlecht": "w", "massnahme_zeitraum": "18.05.2026 - 23.08.2026"})
    pruefe("Improfy-Block trägt den Maßnahme-Zeitraum",
           block["zeitraum"] == "18.05.2026 - 23.08.2026" and block["jobtitel"] == "Teilnehmerin", block)
    pruefe("Mehrere Sprachen in einer Zelle werden getrennt",
           [e["sprache"] for e in LB._sprachen({"sprache": "Deutsch/Paschtu"})] == ["DEUTSCH", "PASCHTU"],
           LB._sprachen({"sprache": "Deutsch/Paschtu"}))
    pruefe("Ohne Zeitraum bleibt es bei aktuell",
           cv_pdf.improfy_block({})["zeitraum"] == "aktuell")
    # Seit dem Umbau auf die Sammlung teilen sich alle Nummern ein Gerüst. Geprüft wird
    # deshalb nicht mehr „jede Vorlage hat ihre eigene Datei", sondern: jede angebotene
    # Nummer führt auf eine vorhandene Datei **und** hat eine eigene Beschreibung.
    # Ohne das Zweite stünden im Auswahlfeld 23 Einträge, die alle gleich aussähen.
    import cv_sammlung
    pruefe("Jede angebotene Vorlage führt auf eine vorhandene Datei",
           all(os.path.exists(os.path.join("templates", cv_pdf.DESIGN_DATEI[sch]))
               for sch, _, _ in cv_pdf.DESIGNS),
           [sch for sch, _, _ in cv_pdf.DESIGNS
            if not os.path.exists(os.path.join("templates", cv_pdf.DESIGN_DATEI.get(sch, "")))])
    _skins = [cv_sammlung.skin(sch) for sch, _, _ in cv_pdf.DESIGNS]
    pruefe("Keine zwei Vorlagen sehen gleich aus",
           len({(k["kopfform"], k["farbe"], k["farbe2"], k["bewertung"]) for k in _skins})
           == len(_skins),
           f"{len(_skins)} Vorlagen")
    pruefe("Jede Vorlage nennt eine Kopfform, die das Gerüst kennt",
           all(k["kopfform"] in ("band", "links", "ecke", "flaeche", "rahmen", "schlicht")
               and k["bewertung"] in ("sterne", "punkte", "balken", "text") for k in _skins),
           [k["kennung"] for k in _skins
            if k["kopfform"] not in ("band", "links", "ecke", "flaeche", "rahmen", "schlicht")])
    with A.app.test_request_context():
        html = cv_pdf.html_bauen(flask_render,
                                 dict(daten, massnahme_zeitraum="01.01.2026 - 01.03.2026"),
                                 cv_pdf.DESIGNS[0][0])
    # Seit dem Umbau auf den gemeinsamen Rahmen gibt es keine festen Seiten mehr –
    # der Inhalt fließt über so viele, wie er braucht. Geprüft wird deshalb, dass die
    # Bausteine da sind, nicht wie viele Seiten herauskommen.
    pruefe("Die erste Vorlage rendert Name, Zeitraum und die Kapitel",
           daten["nachname"] in html and "01.01.2026 - 01.03.2026" in html
           and html.count('class="kapitel"') >= 2,
           f"{html.count('class=\"kapitel\"')} Kapitel")
    # Der eigentliche Fehler war eine **feste Seitenhöhe**: Was darüber hinausging, war weg.
    # `overflow: hidden` an Zierelementen (Bildkopf, Linie neben der Überschrift) ist in
    # Ordnung – geprüft wird deshalb nur die feste Höhe, und Kommentare zählen nicht mit.
    schlecht = [os.path.basename(f) for f in vorlagen_dateien()
                if re.search(r"height:\s*29[0-9]mm", ohne_kommentare(offen(f)))]
    pruefe("Keine Vorlage presst den Inhalt in eine feste Seitenhöhe",
           not schlecht, schlecht)

    print("\n8. Unvollständige Eingabe bricht nicht ab")
    r = c.post(f"/kunde/{kid}/lebenslauf", data={"vorname": "Nur", "nachname": "Name"})
    pruefe("Excel entsteht auch ohne Pflichtfelder, Lücken werden gemeldet",
           r.status_code == 200 and int(r.headers.get("X-Fehlende-Felder", 0)) > 0,
           r.headers.get("X-Fehlende-Felder") + " fehlende Felder")

    print("\n9. Bewerbungsfoto")
    import io as _io
    import sys as _sys
    import fotos
    fotos.init()
    # **Ohne Pillow wird gar nichts hinterlegt.** Der Rueckfall im ImportError-Zweig von
    # `verkleinern` oeffnete keine Datei: Alles, was `bildtyp` an den ersten Bytes
    # erkennt, kam durch - auch `image/heic`, das weder Browser noch PDF-Druck
    # darstellen. Ein frischer Checkout ohne `pip install` haette damit bei
    # drei belegten Plaetzen ein echtes Bewerbungsfoto durch einen unlesbaren HEIC-Block
    # ersetzt und „Foto hinterlegt" gemeldet. Geprueft wird das **immer**, nicht nur auf
    # einem Rechner ohne Pillow: `None` in `sys.modules` laesst `from PIL import Image`
    # genau so scheitern wie eine fehlende Installation.
    _heic_huelle = b"\x00\x00\x00\x18ftypheic\x00\x00\x00\x00heicmif1" + b"\x00" * 64
    _pil_merk = _sys.modules.get("PIL", "fehlt")
    _sys.modules["PIL"] = None
    try:
        fotos.verkleinern(_heic_huelle)
        _ohne_pillow = ""
    except fotos.KeinBild as _e:
        _ohne_pillow = str(_e)
    finally:
        if _pil_merk == "fehlt":
            _sys.modules.pop("PIL", None)
        else:
            _sys.modules["PIL"] = _pil_merk
    pruefe("Fehlt Pillow, wird nichts hinterlegt – auch kein erkanntes Format",
           "Pillow" in _ohne_pillow, _ohne_pillow[:70] or "durchgelassen")
    try:
        from PIL import Image
    except ImportError:
        Image = None
    if Image is None:
        pruefe("Pillow fehlt – die Reihen mit erzeugten Bildern entfallen", True,
               "ohne Pillow lässt sich kein Probebild erzeugen")
    else:
        bild = Image.new("RGB", (2000, 2500), (150, 170, 120))
        puffer = _io.BytesIO(); bild.save(puffer, "JPEG", quality=95)
        gross = puffer.getvalue()
        r = c.post(f"/kunde/{kid}/foto", data={"foto": (_io.BytesIO(gross), "Probe.jpg")},
                   content_type="multipart/form-data")
        f = fotos.foto(kid)
        pruefe("Foto wird gespeichert und auf Druckgröße verkleinert",
               r.status_code == 200 and f and f["bytes"] < len(gross) and f["breite"] <= 1200,
               f"{len(gross)//1024} kB → {f['bytes']//1024} kB, {f['breite']}×{f['hoehe']}")
        pruefe("Foto lässt sich wieder ausliefern",
               c.get(f"/kunde/{kid}/foto.jpg").status_code == 200)
        # Seit mehreren Fotos je Kunde steht nicht mehr „Foto hinterlegt" da, sondern
        # eine Karte je Bild. Geprueft wird deshalb, dass die Bildadresse des Kunden
        # im Formular steht - das ist die Aussage, auf die es ankommt.
        _f = fotos.alle_fotos(kid)
        _seite = c.get(f"/kunde/{kid}/lebenslauf").get_data(as_text=True)
        pruefe("Formular zeigt das hinterlegte Foto",
               bool(_f) and f"/kunde/{kid}/foto/{_f[0]['id']}.jpg" in _seite,
               [x["platz"] for x in _f])

        # --- Welchen Knopf drueckt Enter? -------------------------------------
        # Seit die Fotokarten ihre Knoepfe im grossen Formular tragen, stand der
        # Entfernen-Knopf des ersten Bildes ganz vorn - und damit loeschte ein Enter im
        # Vornamen das Bewerbungsfoto des Kunden, endgueltig und ohne Rueckfrage. Der
        # Testclient sieht das nie, er klickt gezielt; im Browser entscheidet die
        # Reihenfolge im Blatt. Geprueft wird deshalb der erste Absende-Knopf selbst.
        _erster = erster_absender(_seite)
        pruefe("Enter im großen Formular baut die Excel und löscht nichts",
               'formaction="/kunde/%d/lebenslauf"' % kid in _erster
               and " name=" not in _erster and "formnovalidate" not in _erster
               and "disabled" not in _erster and "/foto/" not in _erster,
               " ".join(_erster.split()) or "kein Absende-Knopf im Formular gefunden")

        # `hidden` allein versteckt nichts. Das Attribut wirkt ueber die Grundregel des
        # Browsers ([hidden]{display:none}), und die steht im Ursprung des Browsers:
        # jede Regel von uns schlaegt sie. `stil.css` gibt jedem Knopf
        # `display:inline-flex` - damit stand dieser Knopf als leere gruene Pille
        # zwischen „Was die Taskforce schon weiss" und der Karte „Person", mit der Maus
        # anklickbar, mit Tab erreichbar, mit der Leertaste ausloesbar (am laufenden
        # Server gemessen: display inline-flex, 34x18 px). Deshalb wird hier beides
        # geprueft: das Attribut am Knopf **und** die Regel, die es durchsetzt.
        _css = c.get("/static/stil.css").get_data(as_text=True)
        pruefe("Der Standardknopf trägt `hidden` – und die Stilvorlage setzt das durch",
               " hidden" in _erster
               and re.search(r"\[hidden\][^{]*\{[^}]*display\s*:\s*none\s*!important",
                             _css) is not None,
               "[hidden]-Regel in stil.css: "
               + ("vorhanden" if "[hidden]" in _css else "FEHLT"))

        _fid = _f[0]["id"]
        _design = cv_sammlung.alle()[-1]["kennung"]
        if not cv_pdf.bereit():
            # Fotokarten und Vorlagenwahl stehen im Blatt unter {% if chrome %}. Ohne
            # installiertes Chrome/Edge zeigt die Seite sie gar nicht - dann ist an
            # ihnen nichts zu pruefen und auch nichts kaputt. Gesagt wird es trotzdem,
            # statt still zu ueberspringen.
            pruefe("Ohne Chrome stehen Fotokarten und Vorlagenwahl nicht im Blatt – "
                   "die Prüfungen daran entfallen",
                   "fotokarte" not in _seite
                   and 'type="radio" name="design"' not in _seite)
        else:
            # --- Die beiden Knoepfe an der Fotokarte --------------------------
            # Sie stehen mitten im grossen Formular, hingen aber per `form="..."` an zwei
            # leeren Formularen am Blattende. Abgeschickt wird immer nur das Formular, an
            # dem der Knopf haengt - und dort stand nichts: Ein Klick auf „entfernen"
            # kostete den halb geschriebenen Lebenslauf. Geprueft wird deshalb nicht
            # „kein Fehler", sondern dass die ausgefuellten Felder wieder herausfallen.
            pruefe("Foto-Knöpfe hängen am großen Formular",
                   'form="loesch' not in _seite and 'id="loesch' not in _seite
                   and 'form="platz' not in _seite)
            # Der Testclient fuehrt keine Browservalidierung aus, der Coach aber schon:
            # ohne `formnovalidate` blockt der Browser das Entfernen, solange ein
            # Pflichtfeld (Vorname, Nachname, Zielberuf) leer ist. Der Knopf saehe aus
            # wie immer und taete nichts. Dass er ohne Rueckfrage loescht, ist nur so
            # lange vertretbar, wie er **nicht** der Standardknopf ist - deshalb steht
            # hier beides in einer Pruefung.
            _knopf = _seite[_seite.rfind("<button", 0, _seite.find("entfernen</button>")):
                            _seite.find("entfernen</button>")]
            pruefe("Der Entfernen-Knopf nennt sein Ziel, umgeht die Pflichtfeldprüfung "
                   "und ist nicht der Standardknopf",
                   "formnovalidate" in _knopf and "formaction" in _knopf
                   and "/foto/" in _knopf and "/foto/" not in _erster,
                   " ".join(_knopf.split())[:90])

            # --- Der Knopf, der das Foto hinterlegt ---------------------------
            # Er hat sein eigenes Ziel (`/kunde/<id>/foto`), weil diese Seite auch unter
            # `/lebenslauf?kunde=<id>` steht und ein Knopf ohne `formaction` von dort in
            # 405 liefe. `formnovalidate` wie am Entfernen-Knopf: Ein Foto muss sich
            # auch hinterlegen lassen, solange der Zielberuf noch leer ist. Und er darf
            # **nicht** der erste Absende-Knopf sein - sonst legte Enter wieder Fotos ab.
            _ende = _seite.find("Foto hinterlegen</button>")
            _fotoknopf = _seite[_seite.rfind("<button", 0, _ende):_ende] if _ende > 0 else ""
            pruefe("Der Foto-Knopf nennt sein Ziel, umgeht die Pflichtfeldprüfung und "
                   "ist nicht der Standardknopf",
                   _ende > 0 and 'formaction="/kunde/%d/foto"' % kid in _fotoknopf
                   and "formnovalidate" in _fotoknopf
                   and _seite.find(_erster) < _ende,
                   " ".join(_fotoknopf.split())[:90] or "kein Foto-Knopf im Blatt")

            # Die Vorlagenwahl steht im grossen Formular; „Unterlagen auslesen" ist ein
            # eigenes, vorher geschlossenes Formular. Von dort kam nie ein Design mit -
            # der Coach waehlte Nr. 17, legte den alten Lebenslauf ab und bekam die
            # Seite mit Nr. 1 zurueck. Jetzt faehrt die Kennung als verstecktes Feld mit.
            pruefe("Das Auslese-Formular führt ein Feld für die Vorlage mit",
                   'name="design"' in _seite[:_seite.find("Unterlagen auslesen")])
            _r = c.post(f"/kunde/{kid}/lebenslauf/lesen",
                        data={"rohtext": "Lagerhelfer bei der Testfirma GmbH, Köln",
                              "design": _design})
            _t = _r.get_data(as_text=True)
            pruefe("Gewähltes Design überlebt das Auslesen der Unterlagen",
                   _r.status_code == 200 and 'value="%s" checked' % _design in _t, _design)

            # Auch der leere Zweig reicht die Vorlage durch: Faellt aus den Unterlagen
            # nichts heraus, kommt die Seite mit einer Meldung zurueck - und sprang
            # frueher dabei auf Nr. 1. Bisher fasste diese Stelle kein Test an.
            _r = c.post(f"/kunde/{kid}/lebenslauf/lesen",
                        data={"rohtext": "", "design": _design})
            _t = _r.get_data(as_text=True)
            pruefe("Gewähltes Design überlebt eine leere Auslese",
                   _r.status_code == 200 and "nichts auslesen" in _t
                   and 'value="%s" checked' % _design in _t, _design)

            # Und die beiden Fehlerzweige. Damit sie ueberhaupt vorkommen, wird der
            # Bauschritt absichtlich zum Stolpern gebracht und gleich wieder eingehaengt
            # - sonst bleibt der Zweig fuer immer ungeprueft, und die Vorlage ginge
            # genau dann verloren, wenn der Coach ohnehin von vorn anfangen muss.
            def _stolpern(*_a, **_k):
                raise RuntimeError("Probe: Bau abgebrochen")

            _echt = A.LB.bauen
            A.LB.bauen = _stolpern
            try:
                _t = c.post(f"/kunde/{kid}/lebenslauf",
                            data=dict(daten, design=_design)).get_data(as_text=True)
            finally:
                A.LB.bauen = _echt
            pruefe("Gewähltes Design überlebt einen Fehler im Excel-Bau",
                   "Probe: Bau abgebrochen" in _t
                   and 'value="%s" checked' % _design in _t, _design)

            _echt = A.cv_pdf.bauen
            A.cv_pdf.bauen = _stolpern
            try:
                _t = c.post(f"/kunde/{kid}/lebenslauf/pdf",
                            data=dict(daten, design=_design)).get_data(as_text=True)
            finally:
                A.cv_pdf.bauen = _echt
            pruefe("Gewähltes Design überlebt einen Fehler im PDF-Druck",
                   "PDF nicht erzeugt" in _t
                   and 'value="%s" checked' % _design in _t, _design)

            # Der Platzwechsel wird ueber den zweiten Einstieg geprueft:
            # `/lebenslauf?kunde=` ist der normale Weg und kennt kein POST. Ein Formular
            # ohne gesetztes Ziel liefe von dort in 405 - der Klick verpuffte. Das Ziel
            # kommt deshalb aus der gelieferten Seite, so wie der Browser es macht.
            _seite2 = c.get(f"/lebenslauf?kunde={kid}").get_data(as_text=True)
            _ziel = re.search('name="platz_%d"[^>]*data-ziel="([^"]+)"' % _fid, _seite2)
            _r = c.post(_ziel.group(1),
                        data=dict(daten, design=_design, **{f"platz_{_fid}": "neben1"})) \
                if _ziel else None
            # **Der Titel nennt, was wirklich gemessen wird.** „behält das ausgefüllte
            # Formular" versprach mehr, als die Prüfung hielt: `LB.aus_formular`
            # verwirft **Teilzeilen** (lebenslauf_bauen.py:273-291) - eine Berufsstation
            # mit nur Tätigkeiten, eine Bildungszeile mit nur einer Note, ein Niveau
            # ohne Sprache, Sterne ohne Programm. Der Coach, der in Station 3 erst die
            # Tätigkeiten tippt und dann auf eine Fotokarte klickt, findet den Text
            # nicht wieder. Deshalb wird hier eine solche Teilzeile mitgeschickt und
            # getrennt gemessen, statt sie unter einer zu weiten Überschrift zu
            # verstecken. Geändert wird das Verhalten in dieser Runde nicht - das ist
            # eine eigene Entscheidung, siehe Bericht.
            _teilzeile = "Palettieren im Hochregal"
            _mit_teil = dict(daten, design=_design, **{f"platz_{_fid}": "neben1"})
            for _feld in ("beruf_zeitraum", "beruf_firma", "beruf_jobtitel"):
                _mit_teil[_feld] = list(daten[_feld]) + [""]
            _mit_teil["beruf_taet"] = list(daten["beruf_taet"]) + [_teilzeile]
            _r = c.post(_ziel.group(1), data=_mit_teil) if _ziel else None
            _t = _r.get_data(as_text=True) if _r else ""
            pruefe("Platzwechsel behält die ausgefüllten Felder und vollständige Zeilen",
                   bool(_ziel) and _r.status_code == 200
                   and 'name="angestrebter_job" value="Lagerhelfer"' in _t
                   and daten["ueber_mich"] in _t
                   and 'name="beruf_firma" value="Testfirma GmbH, Köln"' in _t,
                   f"{_ziel.group(1) if _ziel else 'kein Ziel in der Seite'} · "
                   f"{_r.status_code if _r else '-'}")
            # Die Gegenprobe, damit die Lücke gemessen dasteht und nicht behauptet:
            # Wird das Verhalten eines Tages geändert, wird diese Zeile rot und zwingt
            # dazu, auch den Titel darüber wieder anzufassen.
            pruefe("Bekannt und offen: eine Berufszeile aus nur Tätigkeiten überlebt "
                   "den Fotoklick nicht",
                   bool(_ziel) and _teilzeile not in _t,
                   "im Blatt wiedergefunden" if _teilzeile in _t else "verworfen")
            # Die Vorlagenwahl stand fest auf Nr. 1: jedes Neuzeichnen warf die Wahl weg.
            pruefe("Gewähltes Design überlebt den Fotowechsel",
                   'value="%s" checked' % _design in _t, _design)

            # Ein im Dateifeld gewaehltes Bild darf beim Foto-Klick nicht mitgespeichert
            # werden: `fotos.freier_platz` gibt bei drei belegten Plaetzen wieder den
            # Kopf zurueck - das neue Bild ueberschriebe also das Foto von Seite 1, und
            # gemeldet wuerde nur „Foto entfernt.". Geprueft wird: es kommt keines dazu,
            # es aendert sich keines, und der Coach erfaehrt, dass seine Auswahl
            # liegengeblieben ist.
            _stand = {(z["id"], z["platz"], z["bytes"]) for z in fotos.alle_fotos(kid)}
            _r = c.post(f"/kunde/{kid}/foto/{_fid}",
                        data=dict(daten, foto=(_io.BytesIO(gross), "Zweites.jpg"),
                                  **{f"platz_{_fid}": "neben1"}),
                        content_type="multipart/form-data")
            _t = _r.get_data(as_text=True)
            pruefe("Der Foto-Klick legt kein neues Bild ab und sagt das auch",
                   _r.status_code == 200 and "nicht hinterlegt" in _t
                   and {(z["id"], z["platz"], z["bytes"])
                        for z in fotos.alle_fotos(kid)} == _stand,
                   [(z["platz"], z["bytes"]) for z in fotos.alle_fotos(kid)])

            _r = c.post(f"/kunde/{kid}/foto/{_fid}", data=dict(daten, was="loeschen"))
            _t = _r.get_data(as_text=True)
            pruefe("Fotowechsel behält das ausgefüllte Formular",
                   _r.status_code == 200
                   and 'name="angestrebter_job" value="Lagerhelfer"' in _t
                   and daten["ueber_mich"] in _t
                   and 'name="beruf_firma" value="Testfirma GmbH, Köln"' in _t,
                   _r.status_code)
            # Zweiter Tab, Zurueck-Taste, zweimal geklickt: Dann ist das Foto schon fort.
            # „Foto entfernt." waere gelogen; der Platzwechsel sagt in derselben Lage
            # „Dieses Foto gibt es nicht mehr."
            _t = c.post(f"/kunde/{kid}/foto/{_fid}",
                        data=dict(daten, was="loeschen")).get_data(as_text=True)
            pruefe("Ein Foto, das schon fort ist, wird nicht noch einmal als entfernt "
                   "gemeldet",
                   "Dieses Foto gibt es nicht mehr." in _t and "Foto entfernt." not in _t)

        # --- Ein Weg, ein Foto zu hinterlegen -------------------------------
        # Vorher waren es drei halbe: der Excel-Knopf, der PDF-Knopf und diese Route.
        # Weil der versteckte Standardknopf des Blattes auf den Excel-Zweig zeigt, legte
        # **jedes** Enter in einem Textfeld ein weiteres Foto ab, solange im Dateifeld
        # eine Datei stand: drei Enter fuellten alle drei Plaetze, das vierte
        # ueberschrieb still das Kopffoto. Der Coach sah davon nichts - die Antwort ist
        # ein Download. Dreimal auf „Excel erzeugen" tat dasselbe.
        fotos.loeschen(kid)
        _klein = _io.BytesIO()
        Image.new("RGB", (600, 800), (40, 60, 120)).save(_klein, "JPEG", quality=80)
        _klein = _klein.getvalue()
        _r = c.post(f"/kunde/{kid}/foto",
                    data=dict(daten, was="hinterlegen",
                              foto=(_io.BytesIO(gross), "Fotoknopf.jpg")),
                    content_type="multipart/form-data")
        _t = _r.get_data(as_text=True)
        pruefe("Der Foto-Knopf hinterlegt das Bild und behält das ausgefüllte Formular",
               _r.status_code == 200 and "Foto hinterlegt" in _t
               and fotos.foto(kid) is not None
               and 'name="angestrebter_job" value="Lagerhelfer"' in _t
               and daten["ueber_mich"] in _t,
               f"Status {_r.status_code}, {len(fotos.alle_fotos(kid))} Foto(s)")
        pruefe("Ohne gewählte Datei sagt der Foto-Knopf, was fehlt",
               "Keine Datei gewählt." in c.post(f"/kunde/{kid}/foto", data=dict(daten))
               .get_data(as_text=True))
        # Erst mit allen drei Plaetzen belegt zeigt sich das stille Ueberschreiben:
        # `fotos.freier_platz` gibt dann wieder den Kopf zurueck.
        for _name in ("Zweites.jpg", "Drittes.jpg"):
            c.post(f"/kunde/{kid}/foto",
                   data=dict(daten, foto=(_io.BytesIO(gross), _name)),
                   content_type="multipart/form-data")
        _stand = [(z["id"], z["platz"], z["bytes"]) for z in fotos.alle_fotos(kid)]
        pruefe("Drei Fotoplätze lassen sich über den Foto-Knopf belegen",
               len(_stand) == 3, [(z[1], z[2]) for z in _stand])

        # **Auch die Fotokarte darf das Blatt nicht leeren.** `lebenslauf_foto` schuetzt
        # sich mit `if "vorname" in request.form` und zeichnet sonst den gespeicherten
        # Stand; `kunde_foto_aendern` reichte `LB.aus_formular` ungeprueft durch, und
        # ein Aufruf ohne Formular kam mit leeren Feldern zurueck - obwohl der Docstring
        # Gleichlauf mit `lebenslauf_foto` versprach. Der Platz, auf den gelegt wird,
        # ist der, auf dem das Foto schon liegt: `platz_setzen` kehrt dann sofort um,
        # am Bestand aendert sich nichts, und gemessen wird allein das gezeichnete Blatt.
        def _feld(_text, _name):
            _m = re.search(r'name="%s" value="([^"]*)"' % _name, _text)
            return _m.group(1) if _m else None

        _geladen = c.get("/lebenslauf?kunde=%d" % kid).get_data(as_text=True)
        _foto1 = fotos.alle_fotos(kid)[0]
        _r = c.post(f"/kunde/{kid}/foto/{_foto1['id']}",
                    data={"platz_%d" % _foto1["id"]: _foto1["platz"]})
        _t = _r.get_data(as_text=True)
        pruefe("Ein Aufruf ohne Blatt zeigt den gespeicherten Stand, keine leeren Felder",
               _r.status_code == 200 and _feld(_geladen, "vorname")
               and _feld(_t, "vorname") == _feld(_geladen, "vorname")
               and [(z["id"], z["platz"], z["bytes"])
                    for z in fotos.alle_fotos(kid)] == _stand,
               "Vorname im Blatt: %r · gespeichert: %r"
               % (_feld(_t, "vorname"), _feld(_geladen, "vorname")))

        # --- Was kein Bild ist, kommt nicht herein ---------------------------
        # Der schlimmste Fall, den dieses Formular hatte: `fotos.verkleinern` fing jede
        # Ausnahme und behielt „lieber das Original als gar kein Foto" - mit
        # aufgestempeltem `image/jpeg`. Zusammen mit `freier_platz` (bei drei belegten
        # Plaetzen wieder der Kopf) und `ON CONFLICT ... DO UPDATE` ersetzte eine
        # 17-Byte-Textdatei das echte Bewerbungsfoto, gemeldet wurde „Foto hinterlegt,
        # 0 kB.". Kein Laborfall: `BILDENDUNGEN` fuehrt `.heic`, und das hier
        # installierte Pillow kennt HEIC nicht - ein iPhone-Foto aus „CVs Köln" lief
        # genau in diesen Zweig. Geprueft wird deshalb dreierlei: Es wird abgelehnt,
        # der Grund steht da, und der belegte Platz ist unberuehrt.
        _r = c.post(f"/kunde/{kid}/foto",
                    data=dict(daten, foto=(_io.BytesIO(b"das ist kein Bild"),
                                           "kein_bild.txt")),
                    content_type="multipart/form-data")
        _t = _r.get_data(as_text=True)
        pruefe("Eine Datei, die kein Bild ist, wird abgelehnt und überschreibt nichts",
               _r.status_code == 200 and "kein lesbares Bild" in _t
               and "Foto hinterlegt" not in _t
               and [(z["id"], z["platz"], z["bytes"])
                    for z in fotos.alle_fotos(kid)] == _stand,
               [(z["platz"], z["dateiname"], z["bytes"]) for z in fotos.alle_fotos(kid)])
        # HEIC bekommt eine eigene Auskunft: „ging nicht" sagt dem Coach nicht, was er
        # tun soll. Die Bytes sind eine ISO-BMFF-Huelle mit der Marke `heic` - genau
        # das, was ein iPhone liefert.
        _heic = (b"\x00\x00\x00\x18ftypheic\x00\x00\x00\x00heicmif1"
                 + b"\x00" * 64)
        _r = c.post(f"/kunde/{kid}/foto",
                    data=dict(daten, foto=(_io.BytesIO(_heic), "IMG_4711.heic")),
                    content_type="multipart/form-data")
        _t = _r.get_data(as_text=True)
        pruefe("Ein HEIC vom iPhone wird abgelehnt und sagt, was zu tun ist",
               _r.status_code == 200 and "HEIC" in _t and "Maximal kompatibel" in _t
               and [(z["id"], z["platz"], z["bytes"])
                    for z in fotos.alle_fotos(kid)] == _stand,
               fotos.bildtyp(_heic))
        # **AVIF ist kein Sonderfall mehr.** Dieses Pillow liest es (gemessen:
        # `features.check("avif")` ist wahr), und Chrome - ueber das der PDF-Druck
        # laeuft - zeigt es seit Fassung 85; AVIF steht deshalb in `fotos.FORMATE` und
        # kommt weiter unten in der Formatreihe wirklich durch `verkleinern`. Hier
        # steht der andere Fall: eine AVIF-Huelle ohne Bildinhalt, wie sie auf einem
        # Rechner mit AVIF-losem Pillow auch von einem echten Bild uebrig bliebe. Sie
        # darf **nicht** den Satz „bitte noch einmal aus dem Chat herunterladen"
        # bekommen - ein zweiter Download aendert nichts an einer Bildbibliothek, und
        # der Coach laedt endlos neu.
        _avif = (b"\x00\x00\x00\x18ftypavif\x00\x00\x00\x00avifmif1"
                 + b"\x00" * 64)
        _r = c.post(f"/kunde/{kid}/foto",
                    data=dict(daten, foto=(_io.BytesIO(_avif), "IMG_4712.avif")),
                    content_type="multipart/form-data")
        _t = _r.get_data(as_text=True)
        pruefe("Ein unlesbares AVIF schickt den Coach nicht zum Neuladen",
               _r.status_code == 200 and "AVIF" in _t
               and "noch einmal aus dem Chat" not in _t
               and "Maximal kompatibel" not in _t
               and [(z["id"], z["platz"], z["bytes"])
                    for z in fotos.alle_fotos(kid)] == _stand,
               _t[_t.find("AVIF-Bild"):][:60] or "kein AVIF-Rat")
        pruefe("Der Bildtyp kommt aus den ersten Bytes, nicht aus der Endung",
               fotos.bildtyp(_klein) == "image/jpeg"
               and fotos.bildtyp(_heic) == "image/heic"
               and fotos.bildtyp(_avif) == "image/avif"
               and fotos.bildtyp(b"das ist kein Bild") is None)

        # --- Das vierte Foto nennt den Platz, den es ersetzt ------------------
        # Alle drei Plaetze sind belegt, `freier_platz` gibt wieder den Kopf zurueck -
        # das naechste Bild kostet Seite 1. Das darf passieren, aber nicht
        # stillschweigend: Vorher stand daneben nur „Foto hinterlegt, 2 kB", und im
        # gelieferten Blatt kam „ersetzt" kein einziges Mal vor. Im Echtbestand hat
        # genau ein Kunde Fotos, und zwar alle drei Plaetze belegt - der einzige echte
        # Fall steht also exakt in diesem Zustand.
        _vorher = fotos.alle_fotos(kid)[0]
        _r = c.post(f"/kunde/{kid}/foto",
                    data=dict(daten, foto=(_io.BytesIO(_klein), "Viertes.jpg")),
                    content_type="multipart/form-data")
        _t = _r.get_data(as_text=True)
        _nachher = fotos.alle_fotos(kid)
        # Gesucht wird der Satz, nicht nur das Wort „ersetzt": Das steht seit dieser
        # Runde auch im Hinweis neben dem Knopf, und eine Pruefung darauf waere grün,
        # ohne dass die Meldung selbst etwas sagt.
        pruefe("Das vierte Foto nennt den Platz und das Bild, das es ersetzt hat",
               _r.status_code == 200 and "Foto hinterlegt" in _t
               and ("auf Platz „%s“." % fotos.PLAETZE[0][1]) in _t
               and "wurde dabei ersetzt und ist fort" in _t
               and (_vorher["dateiname"] or "") in _t
               and len(_nachher) == 3 and _nachher[0]["dateiname"] == "Viertes.jpg",
               [(z["platz"], z["dateiname"]) for z in _nachher])
        # Und die Seite verspricht daneben nicht mehr das Gegenteil: Solange alle
        # Plaetze belegt sind, heisst das Feld nicht „hinzufügen".
        # Der Text im Blatt ist umgebrochen; gesucht wird deshalb im zusammengezogenen
        # Fliesstext, nicht im Quelltext mit seinen Zeilenenden.
        _seite3 = " ".join(c.get(f"/lebenslauf?kunde={kid}")
                           .get_data(as_text=True).split())
        pruefe("Bei belegten Plätzen sagt das Formular austauschen, nicht hinzufügen",
               "Foto austauschen" in _seite3 and "Weiteres Foto hinzufügen" not in _seite3
               and "Plätze sind belegt" in _seite3,
               "Foto austauschen" in _seite3)
        # Der Ausgangsstand fuer alles, was danach kommt: wieder drei Plaetze belegt,
        # aber mit der neuen Bildnummer auf dem Kopf.
        _stand = [(z["id"], z["platz"], z["bytes"]) for z in fotos.alle_fotos(kid)]

        # Enter in einem Textfeld loest den ersten Absende-Knopf aus; sein Ziel steht in
        # der Seite. Viermal, mit einer Datei im Fotofeld.
        # **Diese Reihe war schon am Stand vor dem Umbau gruen**: `lebenslauf_erzeugen`
        # hat das Feld `foto` nie gelesen, der Enter-Weg war also nie der Weg, ueber den
        # Fotos verschwanden - das war der PDF-Knopf, mehrfach geklickt (die Reihe
        # darunter). Gefaehrlich wurde der Excel-Zweig erst, als Runde 2 dieses Umbaus
        # den Upload dorthin legte, wo der versteckte Standardknopf haengt. Die Reihe
        # bleibt als Waechter genau dafuer stehen: Wer den Upload je wieder an diese
        # Route haengt, macht sie rot.
        _enterziel = re.search('formaction="([^"]+)"', _erster)
        for _ in range(4):
            c.post(_enterziel.group(1) if _enterziel else f"/kunde/{kid}/lebenslauf",
                   data=dict(daten, foto=(_io.BytesIO(_klein), "Enterbild.jpg")),
                   content_type="multipart/form-data")
        pruefe("Vier Enter mit gewählter Datei legen kein Foto ab und überschreiben keines",
               bool(_enterziel)
               and [(z["id"], z["platz"], z["bytes"])
                    for z in fotos.alle_fotos(kid)] == _stand,
               [(z["platz"], z["bytes"]) for z in fotos.alle_fotos(kid)])

        _r = c.post(f"/kunde/{kid}/lebenslauf",
                    data=dict(daten, foto=(_io.BytesIO(_klein), "Excelfoto.jpg")),
                    content_type="multipart/form-data")
        pruefe("Der Excel-Knopf baut die Excel und rührt die Fotos nicht an",
               _r.status_code == 200 and len(_r.data) > 20000
               and [(z["id"], z["platz"], z["bytes"])
                    for z in fotos.alle_fotos(kid)] == _stand,
               f"{len(_r.data)} Bytes, {len(fotos.alle_fotos(kid))} Foto(s)")
        if cv_pdf.bereit():
            _r = c.post(f"/kunde/{kid}/lebenslauf/pdf",
                        data=dict(daten, foto=(_io.BytesIO(_klein), "PDFfoto.jpg")),
                        content_type="multipart/form-data")
            pruefe("Der PDF-Knopf druckt und rührt die Fotos nicht an",
                   _r.status_code == 200
                   and [(z["id"], z["platz"], z["bytes"])
                        for z in fotos.alle_fotos(kid)] == _stand,
                   f"{len(_r.data)} Bytes, {len(fotos.alle_fotos(kid))} Foto(s)")

        # --- Eine Liste erlaubter Formate, nicht drei ------------------------
        # `BILDENDUNGEN` kannte kein `.gif`, die Meldung in `verkleinern` versprach GIF
        # ausdruecklich, und der Einzelknopf nahm eines an: `aufnehmen([("gut.gif", …)])`
        # sagte „keine Bilddatei", derselbe Inhalt ueber den Knopf „Foto hinterlegt,
        # 400 kB". Geprueft wird jetzt gegen die eine Liste `fotos.FORMATE`: Was dort
        # steht, muss durch `verkleinern` kommen und darf von `aufnehmen` nicht am Namen
        # abgewiesen werden.
        _formatfehler = []
        for _name, _endungen in fotos.FORMATE:
            _p = _io.BytesIO()
            try:
                Image.new("RGB", (900, 1200), (120, 140, 160)).save(_p, _name.upper())
            except Exception as _e:
                # Kann dieses Pillow das Format nicht einmal schreiben, kann
                # `verkleinern` es auch nicht lesen - dann verspricht `FORMATE_TEXT`
                # zu viel. Das ist ein Befund und kein Grund, die Reihe zu ueberspringen.
                _formatfehler.append("%s: Pillow kann das hier nicht (%s)" % (_name, _e))
                continue
            _roh = _p.getvalue()
            try:
                fotos.verkleinern(_roh)
            except fotos.KeinBild as _e:
                _formatfehler.append("%s: %s" % (_name, _e))
                continue
            _abgelehnt = fotos.aufnehmen(
                [("kein_kundenname_%s%s" % (_name.lower(), _endungen[0]), _roh)])[2]
            if _abgelehnt:
                _formatfehler.append("aufnehmen: %s" % _abgelehnt[0])
        pruefe("Jedes Format aus der einen Liste kommt durch – auch über den Sammelweg",
               not _formatfehler, _formatfehler or fotos.FORMATE_TEXT)
        pruefe("Und was kein Bild ist, wird auf beiden Wegen mit demselben Satz "
               "abgelehnt",
               fotos.FORMATE_TEXT in (fotos.aufnehmen(
                   [("notiz.txt", b"das ist kein Bild")])[2] or [""])[0],
               fotos.aufnehmen([("notiz.txt", b"das ist kein Bild")])[2])

        # --- Ein unlesbares Bild bricht das Einlesen eines Ordners nicht ab ---
        # Der Alltag: zwanzig Bilder aus „CVs Köln" in einem Ordner, eine IMG_4711.heic
        # dabei. `eingang_ablegen` wirft seit der Sperre `KeinBild`; ungeschuetzt in der
        # Schleife brach damit der ganze Durchgang ab, alles alphabetisch Dahinterliegende
        # erreichte den Eingang nie, und die Seite zeigte nur den HEIC-Rat.
        import shutil as _shutil
        import tempfile as _tempfile
        _bilderordner = _tempfile.mkdtemp(prefix="fototest-")
        try:
            _p = _io.BytesIO()
            Image.new("RGB", (800, 1000), (90, 110, 130)).save(_p, "JPEG")
            for _n, _inhalt in (("a_IMG_0001.jpg", _p.getvalue()),
                                ("b_IMG_4711.heic", _heic_huelle),
                                ("c_IMG_0002.jpg", _p.getvalue())):
                with open(os.path.join(_bilderordner, _n), "wb") as _f:
                    _f.write(_inhalt)
            _eingang_vorher = fotos.stand()["eingang"]
            _r = c.post("/lebenslauf/fotos", data={"was": "ordner",
                                                   "ordner": _bilderordner})
            _t = " ".join(_r.get_data(as_text=True).split())
            pruefe("Ein unlesbares Bild hält die beiden anderen nicht auf",
                   _r.status_code == 200
                   and fotos.stand()["eingang"] == _eingang_vorher + 2
                   and "2 liegen im Eingang" in _t,
                   "%d → %d im Eingang" % (_eingang_vorher, fotos.stand()["eingang"]))
            pruefe("Und der Grund für das unlesbare Bild steht auf der Seite",
                   "b_IMG_4711.heic" in _t and "HEIC" in _t,
                   _t[_t.find("Nicht angenommen"):][:90] or "kein Grund genannt")
        finally:
            _shutil.rmtree(_bilderordner, ignore_errors=True)
        for _z in fotos.eingang():          # den Eingang wieder leeren
            fotos.eingang_verwerfen(_z["id"])

        # --- Gleiche Gruende einmal, und nichts faellt stumm weg -------------
        # Acht iPhone-Bilder in einem Ordner ergaben acht Zeilen mit demselben
        # 200-Zeichen-Satz - ohne Pillow steht er sogar bei jeder Datei, gleich welchen
        # Formats. Zwanzig davon waren nicht mehr zu lesen. Jetzt steht der Grund
        # einmal da, mit allen Namen davor.
        _ordner2 = _tempfile.mkdtemp(prefix="fototest-viele-")
        try:
            for _i in range(8):
                with open(os.path.join(_ordner2, "IMG_%04d.heic" % _i), "wb") as _f:
                    _f.write(_heic_huelle)
            _r = c.post("/lebenslauf/fotos", data={"was": "ordner",
                                                   "ordner": _ordner2})
            _t = " ".join(_r.get_data(as_text=True).split())
            _block = _t[_t.find("Nicht angenommen"):] if "Nicht angenommen" in _t else ""
            pruefe("Acht gleiche Gründe stehen als eine Zeile, mit allen Namen davor",
                   _r.status_code == 200 and _block.count("Maximal kompatibel") == 1
                   and all("IMG_%04d.heic" % _i in _block for _i in range(8)),
                   "%d× derselbe Grund" % _block.count("Maximal kompatibel"))
        finally:
            _shutil.rmtree(_ordner2, ignore_errors=True)
        for _z in fotos.eingang():
            fotos.eingang_verwerfen(_z["id"])
        # Beim Hochladen traegt jeder Grund seinen Dateinamen, es bleiben also acht
        # verschiedene. Dann greift die Kuerzung der Seite: Sechs stehen da - und der
        # Rest wird gezaehlt statt stillschweigend weggelassen. Vorher fielen bei
        # zwanzig Bildern vierzehn Gruende weg, ohne ein Wort darueber.
        _r = c.post("/lebenslauf/fotos",
                    data={"was": "hochladen",
                          "bilder": [(_io.BytesIO(b"das ist kein Bild"),
                                      "notiz_%d.txt" % _i) for _i in range(8)]},
                    content_type="multipart/form-data")
        _t = " ".join(_r.get_data(as_text=True).split())
        pruefe("Was über die sechs gezeigten Gründe hinausgeht, wird gezählt",
               _r.status_code == 200 and "und 2 weitere" in _t,
               _t[_t.find("Nicht angenommen"):][:120] or "nichts abgelehnt")
        for _z in fotos.eingang():
            fotos.eingang_verwerfen(_z["id"])

        # --- Melden, nicht still loeschen ------------------------------------
        # `was=loeschen` raeumt **alle** Fotos des Kunden ab und meldete dafuer „Foto
        # entfernt." in der Einzahl - ohne Rueckfrage und ohne zu sagen, was fort ist.
        # Kein Knopf zeigt heute dorthin, der Weg bleibt aber ansprechbar.
        _vorher_alle = fotos.alle_fotos(kid)
        _r = c.post(f"/kunde/{kid}/foto", data=dict(daten, was="loeschen"))
        _t = " ".join(_r.get_data(as_text=True).split())
        pruefe("Das Sammellöschen sagt, wie viele Fotos es entfernt hat und welche",
               len(_vorher_alle) > 1 and not fotos.alle_fotos(kid)
               and ("%d Fotos entfernt." % len(_vorher_alle)) in _t
               and all((z["dateiname"] or "") in _t for z in _vorher_alle),
               "%d Bilder, Meldung: %s"
               % (len(_vorher_alle), _t[_t.find("Fotos entfernt"):][:60]))
    # Beide Namen sind erfunden – hier standen zwei echte Kunden, und diese Datei ist
    # versioniert. Geprueft wird unveraendert dasselbe: der Dateiname aus Vor- und
    # Nachname mit Unterstrich, und derselbe Name mitten in einem Satz.
    kunden = [{"id": 1, "name": "Nabil Musterbewerber"}, {"id": 2, "name": "Ana Popa"}]
    pruefe("Foto wird über den Dateinamen zugeordnet",
           fotos.zuordnen("Nabil_Musterbewerber.jpg", kunden) == 1
           and fotos.zuordnen("Bild von Ana Popa.png", kunden) == 2)
    pruefe("Unklare Dateinamen werden nicht zugeordnet",
           fotos.zuordnen("IMG_2931.jpg", kunden) is None
           and fotos.zuordnen("unbekannt.png", kunden) is None)
    pruefe("Fotoseite antwortet", c.get("/lebenslauf/fotos").status_code == 200)
    pruefe("Unbekannter Ordner wird gemeldet, nicht geschluckt",
           "gibt es nicht" in c.post("/lebenslauf/fotos",
                                     data={"ordner": "Z:/gibtesnicht"}).get_data(as_text=True))
    c.post(f"/kunde/{kid}/foto", data={"was": "loeschen"})
    pruefe("Foto lässt sich entfernen", fotos.foto(kid) is None)
    pruefe("Kein Trägerlogo auf dem Lebenslauf – das Papier gehört dem Bewerber",
           not any("logo" in ohne_kommentare(offen(f)).lower() for f in vorlagen_dateien()),
           [os.path.basename(f) for f in vorlagen_dateien()
            if "logo" in ohne_kommentare(offen(f)).lower()])
    # Der Umbruchschutz steht in einem der beiden gemeinsamen Rahmen; eine Vorlage,
    # die keinen einbindet, müsste ihn selbst mitbringen.
    pruefe("Jede Vorlage schützt Einträge vor dem Umbruch",
           all("break-inside: avoid" in offen(f) or "_fein.html" in offen(f)
               or "_rahmen.html" in offen(f) for f in vorlagen_dateien()),
           [os.path.basename(f) for f in vorlagen_dateien()
            if "break-inside: avoid" not in offen(f) and "_fein.html" not in offen(f)
            and "_rahmen.html" not in offen(f)])

    print("\n10. Vorlagenwahl und Einstieg")
    import re as _re
    import cv_sammlung as _CS
    _kid = LB.uebersicht()[0]["id"]
    _seite = c.get("/lebenslauf?kunde=%d" % _kid).get_data(as_text=True)
    _angeboten = set(_re.findall('name="design" value="([^"]+)"', _seite))
    _alle = {k["kennung"] for k in _CS.alle()}
    # Der Kern der Sache: was die Sammlung kennt, muss im Formular anklickbar sein.
    # Eine Vorlage, die es gibt, die aber niemand auswaehlen kann, existiert nicht.
    pruefe("Jede Vorlage der Sammlung laesst sich im Formular auswaehlen",
           _alle <= _angeboten, sorted(_alle - _angeboten))
    pruefe("Die Auswahl bietet nichts an, was es nicht gibt",
           _angeboten <= _alle, sorted(_angeboten - _alle))
    # Gewaehlt wird nach Bild. Fehlt eines, waehlt der Coach blind.
    pruefe("Jede Vorlage hat ein eigenes Vorschaubild",
           all(os.path.exists(os.path.join("static", k["eigen"])) for k in _CS.alle()),
           [k["kennung"] for k in _CS.alle()
            if not os.path.exists(os.path.join("static", k["eigen"]))])
    # Das Bild muss aus der Vorlage selbst stammen, sonst waehlt der Coach etwas
    # anderes, als hinterher gedruckt wird.
    pruefe("Die Auswahl zeigt den eigenen Nachbau, nicht das Werbebild des Designers",
           "/static/vorlagen/" in _seite and "vorlagen/nr25.png" in _seite)

    _ohne = c.get("/lebenslauf")
    pruefe("Lebenslauf oeffnet ohne Umweg ueber einen Kunden",
           _ohne.status_code == 200 and "Fu\u0308r wen" in _ohne.get_data(as_text=True)
           or "F\u00fcr wen" in _ohne.get_data(as_text=True),
           _ohne.status_code)
    for _mist in ("abc", "-1", "0", "99999999999999999999", "1%27"):
        _r = c.get("/lebenslauf?kunde=" + _mist)
        pruefe("Unsinnige Kundennummer (%s) fuehrt zur Auswahl, nicht zum Absturz" % _mist,
               _r.status_code == 200, _r.status_code)

    # --- Die Knoepfe, von beiden Einstiegen aus -----------------------------------
    # Die Luecke, die den toten Excel-Knopf monatelang verdeckt hat: `os_test` ruft jede
    # Seite einmal mit GET auf, und dieser Test hat bisher direkt auf die Route gepostet,
    # die die Excel erzeugt. Beides geht am Fehler vorbei - der sass im **Weg vom Knopf
    # zur Route**. Dieselbe Seite ist unter zwei Adressen erreichbar:
    #
    #   /kunde/<id>/lebenslauf   kennt GET und (auf derselben Adresse) POST
    #   /lebenslauf?kunde=<id>   kennt nur GET - und das ist der normale Weg: die
    #                            Kundenauswahl zeigt dorthin, und von dort aus werden
    #                            auch die Fotoknoepfe gedrueckt
    #
    # Das grosse Formular traegt kein `action`, schickt also an die Adresse, auf der man
    # steht. Vom zweiten Einstieg aus lief der Excel-Knopf damit in 405 - der Knopf sah
    # aus wie immer und tat nichts. Geprueft wird deshalb so, wie der Browser es macht:
    # Zieladresse aus dem gelieferten HTML lesen (`formaction` schlaegt `action`, fehlt
    # beides, gilt die aktuelle Adresse) und genau dorthin posten.
    def _knopfziel(seite, adresse, beschriftung):
        stelle = seite.find(beschriftung)
        if stelle < 0:
            return None
        knopf = seite[seite.rfind("<button", 0, stelle):stelle]
        eigen = re.search(r'formaction="([^"]+)"', knopf)
        if eigen:
            return eigen.group(1)
        anfang = seite.rfind("<form", 0, seite.rfind("<button", 0, stelle))
        formular = seite[anfang:seite.find(">", anfang) + 1]
        am_formular = re.search(r'\saction="([^"]+)"', formular)
        return am_formular.group(1) if am_formular else adresse

    for _weg in (f"/kunde/{kid}/lebenslauf", f"/lebenslauf?kunde={kid}"):
        _seite_knopf = c.get(_weg).get_data(as_text=True)
        # Excel. Geprueft wird nicht nur „kein Fehler", sondern dass wirklich eine
        # Tabelle zurueckkommt: Eine hilfsbereit gerenderte Formularseite haette auch
        # 200 geliefert - nur eben keine Datei.
        _ziel = _knopfziel(_seite_knopf, _weg, "Excel erzeugen und herunterladen")
        _r = c.post(_ziel, data=daten) if _ziel else None
        pruefe("Excel-Knopf liefert von %s eine Excel-Datei" % _weg,
               bool(_ziel) and _r.status_code == 200
               and "spreadsheetml" in (_r.headers.get("Content-Type") or "")
               and len(_r.data) > 20000,
               f"{_ziel} · {_r.status_code if _r else '-'} · "
               f"{_r.headers.get('Content-Type') if _r else '-'}")
        # PDF. Der Knopf steht nur da, wenn ein Chrome gefunden wurde - ohne Browser
        # gibt es nichts zu druecken.
        if cv_pdf.bereit():
            _ziel_pdf = _knopfziel(_seite_knopf, _weg, "PDF im gewählten Design")
            _p = c.post(_ziel_pdf, data=dict(daten, design=cv_pdf.DESIGNS[0][0])) \
                if _ziel_pdf else None
            _name_pdf = re.search(r'filename="([^"]+)"',
                                  _p.headers.get("Content-Disposition", "")) if _p else None
            pruefe("PDF-Knopf liefert von %s ein PDF" % _weg,
                   bool(_ziel_pdf) and _p.status_code == 200
                   and (_p.headers.get("Content-Type") or "").startswith("application/pdf")
                   and _p.data[:4] == b"%PDF",
                   f"{_ziel_pdf} · {_p.status_code if _p else '-'} · "
                   f"{_p.headers.get('Content-Type') if _p else '-'}")
            # Das gedruckte Blatt ist ein Testartefakt und hat in ausgabe/ nichts zu
            # suchen - die Excel ueberschreibt sich selbst, das PDF nicht.
            if _name_pdf:
                _pfad_pdf = LB.datei(_name_pdf.group(1))
                if _pfad_pdf and os.path.exists(_pfad_pdf):
                    os.remove(_pfad_pdf)

    # Die Seitenleiste hatte 14 gleichwertige Reiter. Was taeglich gebraucht wird, muss
    # obenauf liegen; waechst die Liste wieder, geht genau das verloren.
    _basis = offen(os.path.join("templates", "basis.html"))
    _a, _e = _basis.find("nav_arbeit = ["), _basis.find("] %}", _basis.find("nav_arbeit = ["))
    _reiter = [z for z in _basis[_a:_e].splitlines() if z.strip().startswith("('")]
    pruefe("Hoechstens sieben Reiter in der taeglichen Arbeit",
           _a > 0 and len(_reiter) <= 7,
           len(_reiter) if _a > 0 else "nav_arbeit nicht gefunden")

    print("\n11. Fotoeingang")
    import io as _io
    from PIL import Image as _Image
    import fotos as _F

    def _bild(farbe=(200, 190, 180)):
        b = _io.BytesIO()
        _Image.new("RGB", (600, 800), farbe).save(b, "JPEG")
        return b.getvalue()

    _vorher = _F.stand()["mit"]
    # Ein Name, der zu niemandem passt - genau der Fall aus der Chat-Gruppe.
    _r = c.post("/lebenslauf/fotos", data={
        "was": "hochladen",
        "bilder": (_io.BytesIO(_bild()), "IMG_1234.jpg")},
        content_type="multipart/form-data")
    pruefe("Ein Foto ohne passenden Namen landet im Eingang statt im Nichts",
           _r.status_code == 200 and _F.stand()["eingang"] >= 1, _F.stand())
    _liste = _F.eingang()
    pruefe("Der Eingang zeigt das Bild mit seinem Dateinamen",
           any(x["dateiname"] == "IMG_1234.jpg" for x in _liste),
           [x["dateiname"] for x in _liste])
    _eid = [x["id"] for x in _liste if x["dateiname"] == "IMG_1234.jpg"][0]
    _b = c.get(f"/lebenslauf/eingang/{_eid}.jpg")
    pruefe("Das Bild im Eingang laesst sich ansehen",
           _b.status_code == 200 and len(_b.data) > 500, _b.status_code)

    _kid2 = LB.uebersicht()[0]["id"]
    _r = c.post("/lebenslauf/fotos", data={"was": "zuordnen", "eingang_id": _eid,
                                           "kunde_id": _kid2})
    pruefe("Zuordnen macht daraus das Foto des Kunden",
           _r.status_code == 200 and _F.foto(_kid2) is not None
           and _F.stand()["mit"] == _vorher + 1, _F.stand())
    pruefe("Nach dem Zuordnen ist das Bild aus dem Eingang verschwunden",
           not any(x["id"] == _eid for x in _F.eingang()))

    # Verwerfen muss auch gehen - sonst staut sich der Eingang mit Unbrauchbarem.
    c.post("/lebenslauf/fotos", data={"was": "hochladen",
                                      "bilder": (_io.BytesIO(_bild()), "IMG_9999.jpg")},
           content_type="multipart/form-data")
    _eid2 = _F.eingang()[-1]["id"]
    c.post("/lebenslauf/fotos", data={"was": "verwerfen", "eingang_id": _eid2})
    pruefe("Verwerfen raeumt den Eingang",
           not any(x["id"] == _eid2 for x in _F.eingang()))

    # Was kein Bild ist, darf gar nicht erst hereinkommen.
    _r = c.post("/lebenslauf/fotos", data={
        "was": "hochladen", "bilder": (_io.BytesIO(b"kein bild"), "notiz.txt")},
        content_type="multipart/form-data")
    pruefe("Eine Datei, die kein Bild ist, wird abgelehnt statt abgelegt",
           _r.status_code == 200 and "Nicht angenommen" in _r.get_data(as_text=True))

    _r = c.post("/lebenslauf/fotos", data={"was": "zuordnen", "eingang_id": 999999,
                                           "kunde_id": _kid2})
    pruefe("Ein Bild, das es nicht mehr gibt, meldet das statt abzustuerzen",
           _r.status_code == 200 and "Eingang" in _r.get_data(as_text=True),
           _r.status_code)

    print("\n12. Ein Blatt, eine Kante")
    _skin = offen(os.path.join("templates", "cv_designs", "cv_skin.html"))
    # Seit dem 21.09.2026 traegt jede Druckseite ein Foto, immer an derselben Stelle.
    # Die drei frueheren Pruefungen dieses Abschnitts galten dem alten Aufbau (alle
    # Fotos gestapelt im Kopf) und sind damit hinfaellig. An ihre Stelle treten die
    # drei Zusagen, die der neue Aufbau macht.
    # 1. Ab Seite 2 steht links ein Foto, wo bisher Text lief. Der Textkoerper muss
    #    deshalb um **genau** die Breite der Fotospalte einruecken - eine andere Zahl
    #    hiesse: Foto auf Text oder eine Kante, die von Seite 1 abweicht.
    #    Seit dem 21.09.2026 steht die Breite nur noch an einer Stelle - als Jinja-Wert
    #    `fotospalte`. Geprueft wird deshalb nicht mehr, ob zwei ausgeschriebene Zahlen
    #    zufaellig gleich sind, sondern dass beide Stellen denselben Wert benutzen und
    #    keine eigene Millimeterzahl mehr tragen.
    _spalte = re.search(r"\{%\s*set\s+fotospalte\s*=\s*'(\d+mm)'\s*if[^%]*?"
                        r"else\s*'(\d+mm)'\s*%\}", _skin, re.S)
    _zelle = re.search(r"td\.links\s*\{\s*width:\s*\{\{\s*(\w+)\s*\}\}", _skin, re.S)
    _einzug = re.search(r"\.textspalte\s*\{\s*padding-left:\s*\{\{\s*(\w+)\s*\}\}",
                        _skin, re.S)
    pruefe("Die Textspalte rueckt um genau die Breite der Fotospalte ein",
           bool(_spalte) and bool(_zelle) and bool(_einzug)
           and _zelle.group(1) == _einzug.group(1) == "fotospalte",
           f"{_spalte.groups() if _spalte else '?'} · Zelle "
           f"{_zelle.group(1) if _zelle else '?'} gegen Einzug "
           f"{_einzug.group(1) if _einzug else '?'}")
    # 2. Jedes Foto haengt absolut im Anker und ist um seine Seitenzahl mal Vorschub
    #    nach unten versetzt. Ohne dieses `top` laegen alle auf Seite 1 uebereinander.
    with A.app.test_request_context():
        _mit_fotos = cv_pdf.html_bauen(flask_render, dict(daten), cv_pdf.DESIGNS[0][0],
                                       bilder=["bild-eins", "bild-zwei"], seiten=4)
        _ohne_fotos = cv_pdf.html_bauen(flask_render, dict(daten), cv_pdf.DESIGNS[0][0],
                                        bilder=[], seiten=4)
    pruefe("Jedes Seitenfoto traegt seinen Seitenvorschub",
           _mit_fotos.count('class="seitenfoto"') == 4
           and all(f"top: calc({i} * {cv_pdf.VORSCHUB})" in _mit_fotos for i in range(4)))
    # 3. Nie mehr Fotos als Seiten: Ein Foto unterhalb der letzten Seite verlaengert das
    #    Dokument um genau die Seite, auf der es steht - danach fehlt wieder eines, und
    #    das Blatt waechst bei jedem Durchgang weiter.
    with A.app.test_request_context():
        _eine_seite = cv_pdf.html_bauen(flask_render, dict(daten), cv_pdf.DESIGNS[0][0],
                                        bilder=["bild-eins", "bild-zwei", "bild-drei"],
                                        seiten=1)
    pruefe("Nie mehr als ein Foto je Seite",
           _eine_seite.count('class="seitenfoto"') == 1
           and _ohne_fotos.count('class="seitenfoto"') == 0,
           f"{_eine_seite.count('class=\"seitenfoto\"')} Fotos auf einer Seite")
    # Sind weniger Fotos hinterlegt als Seiten da sind, fangen sie von vorn an
    # (Entscheidung des Nutzers): zwei Fotos auf vier Seiten ergibt 1, 2, 1, 2.
    pruefe("Weniger Fotos als Seiten: die Reihe wiederholt sich",
           re.findall(r'class="seitenfoto" src="(bild-\w+)"', _mit_fotos)
           == ["bild-eins", "bild-zwei", "bild-eins", "bild-zwei"],
           re.findall(r'class="seitenfoto" src="(bild-\w+)"', _mit_fotos))
    # Ohne Fotos bleibt der Satz der bisherige - kein Einzug, keine Spalte.
    pruefe("Ohne Foto rueckt nichts ein",
           'class="textspalte"' in _mit_fotos and 'class="textspalte"' not in _ohne_fotos)
    # Eine leere Tabellenzelle schrumpft beim automatischen Layout - dann wandert die
    # Textkante. Nur `table-layout: fixed` haelt die angegebene Breite.
    pruefe("Die Kopfreihe rechnet mit fester Breite",
           "table-layout: fixed" in _skin)
    # Der hervorgehobene Improfy-Block muss Rahmen **und** Polsterung ausgleichen,
    # sonst steht seine Zeile 2,2 pt weiter rechts als jede andere Station.
    pruefe("Der hervorgehobene Block steht auf der Kante der uebrigen",
           "margin-left: -4.8mm" in _skin and
           "padding: 2.5mm 3mm 2.5mm 4mm" in _skin)
    # Der Waechter von frueher ("Alles im Kopf richtet sich nach der Spaltenbreite") hing
    # an einem echten Fehler vom 18.09.2026: Das Foto trug eine eigene Breite in
    # Millimetern und ragte damit 9 mm ueber seine Zelle hinaus, bis auf 2 pt an den Text.
    # Die Fotos heissen seit dem 21.09. anders, der Fehler waere derselbe - deshalb hier
    # dieselbe Zusage fuer den neuen Aufbau: Anker und Foto nehmen die Spalte, nicht ein
    # eigenes Mass.
    _anker = re.search(r"\.fotoanker\s*\{([^}]*)", _skin)
    _sfoto = re.search(r"\.seitenfoto\s*\{([^}]*)", _skin)
    pruefe("Alles in der Fotospalte richtet sich nach der Spaltenbreite",
           bool(_anker) and bool(_sfoto)
           and "width: 100%" in _anker.group(1) and "width: 100%" in _sfoto.group(1)
           and not re.search(r"width:\s*\d+(\.\d+)?mm",
                             _anker.group(1) + _sfoto.group(1)),
           f"{(_anker.group(1) if _anker else '?')[:60]!r}")

    print("\n13. Zwei Durchgaenge, Rueckfall und Seitenpruefung")
    import seitenpruefung
    # Die Fotoreihe: Auf dem alten Aufrufweg (`foto` + `weitere`) laeuft vorschau_bauen.py,
    # auf dem neuen (`bilder`) die Kundenakte. Beide muessen dieselbe Reihenfolge ergeben,
    # sonst bekommt der eine Weg die Bilder vertauscht auf die Seiten.
    pruefe("Alter und neuer Aufrufweg ergeben dieselbe Fotoreihe",
           cv_pdf._bilderliste(None, "eins", {"neben2": "drei", "neben1": "zwei"})
           == ["eins", "zwei", "drei"]
           and cv_pdf._bilderliste(["eins", "zwei", "drei"]) == ["eins", "zwei", "drei"],
           cv_pdf._bilderliste(None, "eins", {"neben2": "drei", "neben1": "zwei"}))
    # Ein leerer Platz darf keine Luecke in die Reihe reissen: Sonst bekaeme eine Seite
    # ein leeres `src` und bliebe ohne Bild.
    pruefe("Leere Plaetze fallen aus der Reihe",
           cv_pdf._bilderliste(None, None, {"neben1": "", "neben2": "zwei"}) == ["zwei"]
           and cv_pdf._bilderliste([None, "eins", ""]) == ["eins"])
    # Ohne pymupdf ist die Seitenzahl unbekannt. Eine geratene 1 waere schlimmer als keine
    # Antwort - siehe den Rueckfall weiter unten.
    pruefe("Unlesbares PDF: die Seitenzahl ist unbekannt, nicht 1",
           cv_pdf._seiten_zaehlen(b"das ist kein PDF") is None,
           cv_pdf._seiten_zaehlen(b"das ist kein PDF"))

    # Der Zweipass, ohne Chrome zu bemuehen: `pdf_aus_html` schreibt fuer die Dauer des
    # Tests nur mit, was gedruckt werden soll. Genau darauf kommt es an - auf das HTML,
    # nicht auf die Bytes, die Chrome daraus macht.
    _gedruckt = []
    _echt_drucken, _echt_zaehlen = cv_pdf.pdf_aus_html, cv_pdf._seiten_zaehlen
    _pfad_probe = ""
    try:
        cv_pdf.pdf_aus_html = lambda html, zeit=90: (_gedruckt.append(html) or b"%PDF-1.4 ")
        cv_pdf._seiten_zaehlen = lambda pdf: 3
        with A.app.test_request_context():
            _, _, _pfad_probe = cv_pdf.bauen(
                flask_render, dict(daten), cv_pdf.DESIGNS[0][0], foto="bild-eins",
                dateiname="_selbsttest_zweipass.pdf", weitere={"neben1": "bild-zwei"})
        pruefe("Erster Durchgang misst mit einem Foto",
               len(_gedruckt) == 2 and _gedruckt[0].count('class="seitenfoto"') == 1,
               f"{len(_gedruckt)} Durchgaenge")
        pruefe("Zweiter Durchgang setzt ein Foto je gezaehlter Seite",
               _gedruckt[-1].count('class="seitenfoto"') == 3
               and re.findall(r'class="seitenfoto" src="(bild-\w+)"', _gedruckt[-1])
               == ["bild-eins", "bild-zwei", "bild-eins"],
               re.findall(r'class="seitenfoto" src="(bild-\w+)"', _gedruckt[-1]))
        # Der Rueckfall: Ist die Seitenzahl unbekannt, traegt nur Seite 1 ein Foto - dann
        # darf auch nicht eingerueckt werden, sonst stuende ab Seite 2 eine 61 mm breite
        # leere Spalte neben dem Text.
        _gedruckt.clear()
        cv_pdf._seiten_zaehlen = lambda pdf: None
        with A.app.test_request_context():
            _, _, _pfad_probe = cv_pdf.bauen(
                flask_render, dict(daten), cv_pdf.DESIGNS[0][0], foto="bild-eins",
                dateiname="_selbsttest_zweipass.pdf", weitere={"neben1": "bild-zwei"})
        _zahl_fotos = _gedruckt[-1].count('class="seitenfoto"')
        _mit_einzug = 'class="textspalte"' in _gedruckt[-1]
        pruefe("Ohne Seitenzahl: ein Foto, kein Einzug - der Rueckfall ohne pymupdf",
               len(_gedruckt) == 2 and _zahl_fotos == 1 and not _mit_einzug,
               f"{_zahl_fotos} Fotos, Einzug {'ja' if _mit_einzug else 'nein'}")
        # Ohne jedes Foto bleibt es bei einem Durchgang - es gibt nichts zu verteilen.
        _gedruckt.clear()
        cv_pdf._seiten_zaehlen = lambda pdf: 3
        with A.app.test_request_context():
            _, _, _pfad_probe = cv_pdf.bauen(
                flask_render, dict(daten), cv_pdf.DESIGNS[0][0],
                dateiname="_selbsttest_zweipass.pdf")
        pruefe("Ohne Foto wird nur einmal gedruckt",
               len(_gedruckt) == 1 and 'class="seitenfoto"' not in _gedruckt[0],
               f"{len(_gedruckt)} Durchgaenge")
    finally:
        cv_pdf.pdf_aus_html, cv_pdf._seiten_zaehlen = _echt_drucken, _echt_zaehlen
        if _pfad_probe and os.path.exists(_pfad_probe):
            os.remove(_pfad_probe)

    # Die Regel „hoechstens ein Foto je Seite" hat noch nie etwas gefunden - eine Regel,
    # die nie angeschlagen hat, ist unbewiesen. Deshalb ein Blatt, das sie finden **muss**:
    # zwei hochkante Bilder nebeneinander auf einer Seite.
    if cv_pdf.bereit():
        _hochkant = cv_pdf.foto_uri(_bild((180, 170, 160)), "image/jpeg")
        _stelle = ("<img src='%s' style='position:absolute;left:{}mm;top:20mm;"
                   "width:40mm;height:53mm'>" % _hochkant)
        _kopf = "<html><body style='margin:0'>"
        _zwei = cv_pdf.pdf_aus_html(
            _kopf + _stelle.format(20) + _stelle.format(90) + "</body></html>")
        _eins = cv_pdf.pdf_aus_html(_kopf + _stelle.format(20) + "</body></html>")
        _m_zwei = seitenpruefung.pruefe(_zwei)
        _m_eins = seitenpruefung.pruefe(_eins)
        pruefe("Zwei Fotos auf einer Seite werden gefunden",
               any("2 Fotos auf einer Seite" in m for m in _m_zwei), _m_zwei)
        pruefe("Ein Foto auf einer Seite gilt nicht als Fehler",
               not any("Fotos auf einer Seite" in m for m in _m_eins), _m_eins)
        # Und einmal der ganze Weg mit Chrome: `bauen` wird sonst von keiner Pruefung
        # aufgerufen, obwohl jeder Lebenslauf darueber entsteht.
        #
        # Gedruckt wird mit dem **Lebenslauf-Datensatz** aus `cv_probe`, nicht mit dem
        # Formular-Woerterbuch `daten` von weiter oben. Das Formular spricht
        # `beruf_zeitraum`, `spr_sprache`, `fs_vorhanden`; die Vorlage liest
        # `berufserfahrung`, `sprachen`, `fuehrerschein` und fand davon fast nichts.
        # Ergebnis war **eine** Seite - damit lief der zweite Durchgang nie, und die
        # ganze Mechanik „ein Foto je Seite" wurde von keinem Selbsttest ausgefuehrt.
        import cv_probe
        _pfad_echt = ""
        try:
            with A.app.test_request_context():
                _roh, _name_echt, _pfad_echt = cv_pdf.bauen(
                    flask_render, dict(cv_probe.FAELLE["viel"]), cv_pdf.DESIGNS[0][0],
                    foto=_hochkant, dateiname="_selbsttest_echt.pdf")
            pruefe("bauen druckt ein echtes PDF und legt es ab",
                   _roh[:4] == b"%PDF" and os.path.exists(_pfad_echt)
                   and _name_echt == "_selbsttest_echt.pdf",
                   f"{len(_roh) // 1024} kB")
            # Die Kernzusage, am fertigen PDF gemessen und nicht am HTML: mehr als eine
            # Seite - sonst prueft der Durchlauf den einfachen Fall - und auf jeder
            # Seite genau ein Foto.
            _je_seite = _fotos_je_seite(_roh)
            pruefe("Der Probelauf wird wirklich mehrseitig",
                   len(_je_seite) > 1, f"{len(_je_seite)} Seiten")
            pruefe("Jede gedruckte Seite traegt genau ein Foto",
                   len(_je_seite) > 1 and _je_seite == [1] * len(_je_seite), _je_seite)
            _m_echt = seitenpruefung.pruefe(_roh)
            pruefe("Das gedruckte Blatt haelt die Seitenpruefung aus", not _m_echt, _m_echt)
        finally:
            if _pfad_echt and os.path.exists(_pfad_echt):
                os.remove(_pfad_echt)

        # Und die Gegenprobe zur Regel von eben: Ein mehrseitiges Blatt, auf dem ab
        # Seite 2 das Foto fehlt, **muss** auffallen. Genau das war der Fehler der
        # ersten Runde (leerer 61-mm-Streifen ab Seite 2), und er kam durch jede
        # Pruefung - gemessen wurde nur unter den Seiten, die ein Foto hatten.
        # `break-after` und nicht `break-before`: Das Foto haengt absolut und zaehlt
        # nicht zum Fluss - ein `break-before` am ersten Element im Fluss wirft Chrome
        # weg, und es bliebe bei einer Seite. Genau das ist hier einmal passiert.
        _luecke = cv_pdf.pdf_aus_html(
            _kopf + _stelle.format(20)
            + "<div style='break-after:page'>erste Seite mit Foto</div>"
            + "<div>zweite Seite ohne Foto</div></body></html>")
        _m_luecke = seitenpruefung.pruefe(_luecke)
        pruefe("Eine Seite ohne Foto wird gefunden",
               any("Seiten ohne Foto" in m for m in _m_luecke), _m_luecke)
        pruefe("Ein Blatt ganz ohne Fotos bleibt erlaubt",
               not any("Seiten ohne Foto" in m for m in seitenpruefung.pruefe(
                   cv_pdf.pdf_aus_html(_kopf + "<div>ohne Bild</div></body></html>"))))

        # --- Der Massstab: staucht Chrome das Blatt? ----------------------------------
        # Keine der Regeln in `seitenpruefung` misst die **Groesse**: Das Blatt bleibt
        # A4, auch wenn Chrome den Inhalt auf 92 % verkleinert, weil irgendein Element
        # ueber die Blattkante ragt. Genau so lief der Fehler der Kopfform `ecke` ein
        # halbes Jahr durch - dort war das Foto 48,01 mm breit, in allen anderen
        # Vorlagen 52,12 mm, und alles war gruen.
        # Gemessen wird deshalb die gedruckte Fotobreite, und zwar gegen zwei Dinge:
        #   1. gegen die Fotospalte, wie sie im Aufbau steht (61 mm minus 9 mm Rinne),
        #   2. gegen die uebrigen Vorlagen - staucht Chrome nur eine, faellt sie hier auf.
        # `seitenpruefung` kann das nicht: Sie sieht immer nur ein einzelnes PDF.
        # Je Kopfform genuegt eine Vorlage - die Geometrie des Kopfes kommt aus ihr,
        # nicht aus Farbe oder Schrift. Gedruckt wird der duenne Fall, das spart Zeit.
        import cv_sammlung
        _je_kopfform = {}
        for _kennung, _, _ in cv_pdf.DESIGNS:
            _je_kopfform.setdefault(cv_sammlung.skin(_kennung)["kopfform"], _kennung)
        # Welche Kopfform gibt ihrem Foto einen eigenen Rahmen, und wie stark? Aus dem
        # Aufbau gelesen, nicht hier nachgeschrieben: Wer morgen einer weiteren Vorlage
        # einen Rahmen gibt, bekommt keinen Fehlalarm.
        #
        # Das `\s*` zwischen `%}` und `border:` ist kein Schoenheitsfehler, sondern der
        # Unterschied zwischen einer Meldung, die stimmt, und einer, die in die Irre
        # fuehrt: Rutscht die Eigenschaft beim naechsten Umbruch in die folgende Zeile,
        # fand die alte Regex sie nicht mehr. Die Kopfform waere damit als „ohne Rahmen"
        # gelaufen, in die exakte Pruefung gefallen und dort mit „gestaucht"
        # durchgefallen - an einem Satz, an dem nichts gestaucht ist.
        _rahmen_mm = {f: float(s) for f, s in re.findall(
            r"kopfform == '(\w+)'\s*%\}\s*border:\s*([\d.]+)mm", _skin)}
        _mit_rahmen = set(_rahmen_mm)
        _masse = {}
        for _form, _kennung in _je_kopfform.items():
            with A.app.test_request_context():
                _h = cv_pdf.html_bauen(flask_render, dict(cv_probe.FAELLE["wenig"]),
                                       _kennung, foto=_hochkant, seiten=1)
            _masse[_form] = _fotobreite_mm(cv_pdf.pdf_aus_html(_h))
        # Spaltenbreite und Rinne aus dem Aufbau: ('49', '61', '7', '9'). Gelesen wird,
        # was die Vorlage setzt - so wandert die Erwartung mit, wenn jemand die Spalte
        # aendert, und nur eine echte Stauchung faellt auf.
        _mass = re.search(r"\{%\s*set\s+fotospalte\s*=\s*'(\d+)mm'\s*if[^%]*?"
                          r"else\s*'(\d+)mm'\s*%\}\s*"
                          r"\{%\s*set\s+fotorinne\s*=\s*'(\d+)mm'\s*if[^%]*?"
                          r"else\s*'(\d+)mm'\s*%\}", _skin, re.S)
        _soll = (int(_mass.group(2)) - int(_mass.group(4))) if _mass else 0
        _genau = {f: b for f, b in _masse.items()
                  if f not in _mit_rahmen and f != "links" and b}
        pruefe("Die gedruckte Fotobreite ist die Spalte aus dem Aufbau",
               bool(_mass) and bool(_genau)
               and all(abs(b - _soll) <= 1.0 for b in _genau.values()),
               f"soll {_soll} mm · " + " · ".join(f"{f} {b:.2f}" for f, b in _genau.items())
               # Die gelesenen Rahmen gehoeren in die Meldung: Findet die Regex oben
               # nichts mehr, steht hier `{}` - dann liegt es an ihr und nicht am Satz.
               + f" · Rahmen aus dem Aufbau: {_rahmen_mm}")
        pruefe("Keine Vorlage wird gegenueber den anderen gestaucht",
               len(_genau) > 1 and max(_genau.values()) - min(_genau.values()) <= 0.3,
               f"Spannweite {max(_genau.values()) - min(_genau.values()):.2f} mm"
               if _genau else "nichts gemessen")
        # Die Vorlagen mit eigenem Fotorahmen (und die schmale Randspalte) stehen um die
        # Rahmenstaerke daneben. Frueher stand dafuer ein Fenster von -3,0 bis +0,5 mm,
        # und das war fuer die Randspalte auf 0,42 mm Luft zusammengeschrumpft (gemessen
        # 39,42 gegen eine Untergrenze von 39,0): Ein halber Millimeter mehr Rahmen, und
        # die Pruefung waere rot geworden, ohne dass sich am Satz etwas geaendert haette.
        #
        # Jetzt wird die Rahmenstaerke mitgerechnet, statt sie in einem Puffer zu
        # verstecken. `box-sizing: border-box` (in `_rahmen.html`) legt den Rahmen
        # **innerhalb** der Spalte an, das Bild ist also um zweimal Rahmen schmaler als
        # seine Spalte. Damit hat jede Kopfform einen festen Sollwert und wandert mit,
        # wenn jemand den Rahmen aendert.
        #
        # Toleranz 1,0 mm, und zwar als Messtoleranz und nicht als Puffer: Gemessen am
        # 21.09.2026 weicht keine Kopfform um mehr als 0,41 mm ab (Chrome rundet die
        # Rahmenstaerke auf ganze Geraetepixel, dazu kommt die Rundung des Bildkastens).
        # Eine Stauchung auf 92 % liegt dagegen 3 bis 4 mm daneben und faellt weiter auf.
        _grob = {f: b for f, b in _masse.items() if f in _mit_rahmen or f == "links"}
        _fenster = {f: (int(_mass.group(1)) - int(_mass.group(3)) if f == "links"
                        else _soll) - 2 * _rahmen_mm.get(f, 0.0)
                    for f in _grob} if _mass else {}
        pruefe("Auch Rahmen- und Randspaltenvorlagen bleiben im Mass",
               bool(_fenster) and all(abs(b - _fenster[f]) <= 1.0
                                      for f, b in _grob.items()),
               " · ".join(f"{f} {b:.2f} (soll {_fenster.get(f, 0):.2f})"
                          for f, b in _grob.items()))
        # Die Gegenprobe - eine Regel, die nie angeschlagen hat, ist unbewiesen.
        # Nachgestellt wird der Fehler vom 18.09.2026: ein Zierelement ragt ueber die
        # rechte Blattkante, Chrome verkleinert daraufhin **das ganze Dokument**. Das
        # Blatt bleibt A4, nur der Inhalt schrumpft - genau der Fall, den keine der
        # uebrigen Regeln sieht, weil alle nur nachsehen, ob etwas da ist, nicht wie gross.
        with A.app.test_request_context():
            _h_ecke = cv_pdf.html_bauen(flask_render, dict(cv_probe.FAELLE["wenig"]),
                                        _je_kopfform.get("ecke", cv_pdf.DESIGNS[0][0]),
                                        foto=_hochkant, seiten=1)
        _ueberstand = ("<div style='position:absolute;left:190mm;top:10mm;width:40mm;"
                       "height:40mm;background:#8cc63f;border-radius:50%'></div>")
        _gestaucht = _fotobreite_mm(cv_pdf.pdf_aus_html(
            re.sub(r"(<body[^>]*>)", lambda m: m.group(1) + _ueberstand,
                   _h_ecke, count=1)))
        # `_gestaucht > 0` ist kein Beiwerk: Chrome druckt den abgeschnittenen Zierkreis
        # als haarfeinen Bildblock, und der galt frueher als Foto (hochkant, 0,0 mm
        # breit). Gemessen wurden dann -0,01 mm statt 47,61 mm. Erst seit
        # `seitenpruefung.MINDESTMASS_MM` faellt er heraus - genau das wird hier mit
        # nachgewiesen.
        pruefe("Eine gestauchte Seite wird gefunden",
               _gestaucht > 0 and _soll and abs(_gestaucht - _soll) > 1.0,
               f"{_gestaucht:.2f} mm statt {_soll} mm"
               + (f" ({_gestaucht / _soll:.0%})" if _soll else ""))
    else:
        pruefe("Kein Chrome - der Druckteil von Abschnitt 13 wurde uebersprungen", True)

    # Die Fototoleranz muss mit der Seitenzahl wachsen: Es bleibt eine Restdrift von
    # 0,26 mm je Seite, gegen eine feste Schranke von 3 mm waere ab Seite 13 jedes lange
    # Blatt grundlos rot geworden. Nach oben muss sie unter dem echten Versatz bleiben -
    # der beginnt bei einer halben Spaltenbreite, rund 30 mm.
    pruefe("Die Fototoleranz waechst mit dem Seitenabstand",
           seitenpruefung.ortstoleranz(0) == seitenpruefung.FOTO_ORT_MM
           and seitenpruefung.ortstoleranz(12) > 12 * 0.26
           and seitenpruefung.ortstoleranz(40) > 40 * 0.26
           and seitenpruefung.ortstoleranz(50) < 30,
           f"Seite 13: {seitenpruefung.ortstoleranz(12):.1f} mm · "
           f"Seite 51: {seitenpruefung.ortstoleranz(50):.1f} mm")

    fehl = [n for n, ok in ergebnis if not ok]
    print(f"\n{len(ergebnis) - len(fehl)} von {len(ergebnis)} Prüfungen bestanden.")
    if fehl:
        print("Fehlgeschlagen:", *fehl, sep="\n  - ")
    return 0 if not fehl else 1


if __name__ == "__main__":
    sys.exit(main())
