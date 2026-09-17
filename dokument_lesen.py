# -*- coding: utf-8 -*-
"""Alte Unterlagen ablegen – das Formular füllt sich selbst.

Der Alltag im Coaching: Der Kunde bringt seinen alten Lebenslauf mit. Mal als PDF vom
Jobcenter, mal als Word-Datei, oft als die Improfy-Excel-Vorlage eines früheren
Durchgangs. Niemand soll daraus von Hand vierzig Felder abtippen.

Hier wird die Datei gelesen und in dieselben Felder zerlegt, die auch das Formular
kennt. Drei Wege, in dieser Reihenfolge der Verlässlichkeit:

1. **Improfy-Excel** (`.xlsx`): Die Vorlage hat feste Zellen (`cv/fill_cv.py`), also
   wird sie Zelle für Zelle zurückgelesen. Das ist keine Schätzung, das ist exakt –
   Zeiträume, Tätigkeiten, Sterne, alles kommt vollständig zurück.
2. **PDF, Word, OpenOffice, Text**: Der Text wird herausgeholt und vom Regelwerk in
   `cv_lesen.py` zerlegt. Bei Improfy-Lebensläufen greift das sehr gut, bei fremd
   aufgebauten weniger.
3. **Bilder und Scans ohne Textebene**: Hier ist Schluss. Ein PDF, das nur aus einem
   Foto besteht, hat keinen Text, den man lesen könnte. Statt zu raten sagt das OS das
   klar an – Texterkennung wäre ein eigenes Werkzeug und würde Fehler einschleusen,
   die man im fertigen Lebenslauf nicht mehr sieht.

Mehrere Dateien auf einmal sind erlaubt: Lebenslauf, Zeugnis, alter Excel-Bogen. Was
die verlässlichere Quelle liefert, gewinnt; der Rest füllt nur die Lücken. Nichts wird
überschrieben, was schon steht.
"""
import io
import os
import re

import cv_lesen

# Was wir lesen können, und wie es dem Nutzer gegenüber heißt.
ENDUNGEN = {
    ".xlsx": "Improfy-Excel", ".xlsm": "Improfy-Excel",
    ".pdf": "PDF", ".docx": "Word", ".odt": "OpenOffice", ".rtf": "Text",
    ".txt": "Text", ".md": "Text", ".csv": "Text",
}
BILD = (".jpg", ".jpeg", ".png", ".gif", ".webp", ".heic", ".bmp", ".tif", ".tiff")
# Unter so wenig Text steckt hinter einem PDF in aller Regel ein Scan ohne Textebene.
MINDESTTEXT = 80


# --------------------------------------------------------------- Text aus Dateien
def _text_aus_pdf(rohdaten):
    """PyMuPDF liest den Text in Lesereihenfolge; pypdf ist der Rückfall."""
    try:
        import pymupdf
        with pymupdf.open(stream=rohdaten, filetype="pdf") as doc:
            return "\n".join(seite.get_text("text") for seite in doc)
    except Exception:
        pass
    try:
        from pypdf import PdfReader
        leser = PdfReader(io.BytesIO(rohdaten))
        return "\n".join((s.extract_text() or "") for s in leser.pages)
    except Exception:
        return ""


def _text_aus_word(rohdaten):
    """Absätze und Tabellen – viele Lebensläufe stecken in einer Word-Tabelle."""
    try:
        import docx
        d = docx.Document(io.BytesIO(rohdaten))
    except Exception:
        return ""
    teile = [p.text for p in d.paragraphs]
    for tabelle in d.tables:
        for zeile in tabelle.rows:
            zellen = [z.text.strip() for z in zeile.cells]
            # Doppelte Zellen entstehen durch verbundene Spalten
            entdoppelt = [z for i, z in enumerate(zellen) if i == 0 or z != zellen[i - 1]]
            teile.append("  ".join(z for z in entdoppelt if z))
    return "\n".join(teile)


def _text_aus_odt(rohdaten):
    try:
        from odf import teletype, text as odf_text
        from odf.opendocument import load
        d = load(io.BytesIO(rohdaten))
        return "\n".join(teletype.extractText(p)
                         for p in d.getElementsByType(odf_text.P))
    except Exception:
        return ""


def _text_aus_rtf(rohdaten):
    """Ohne Zusatzpaket: Steuerworte und Klammern entfernen. Reicht für Lebensläufe."""
    roh = rohdaten.decode("latin-1", "replace")
    roh = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: chr(int(m.group(1), 16)), roh)
    roh = re.sub(r"\\par[d]?\b", "\n", roh)
    roh = re.sub(r"\\[a-zA-Z]+-?\d*\s?", " ", roh)
    return re.sub(r"[{}]", "", roh)


def _text_aus_klartext(rohdaten):
    for kodierung in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return rohdaten.decode(kodierung)
        except UnicodeDecodeError:
            continue
    return rohdaten.decode("utf-8", "replace")


def text_aus_datei(dateiname, rohdaten):
    """Reiner Text aus einer Datei. Gibt (text, hinweis) – Hinweis nur, wenn etwas fehlt."""
    endung = os.path.splitext(dateiname or "")[1].lower()
    if endung in BILD:
        return "", f"{dateiname}: Bilder kann das OS nicht lesen – bitte als PDF oder Text ablegen."
    if endung == ".pdf":
        text = _text_aus_pdf(rohdaten)
        if len(text.strip()) < MINDESTTEXT:
            return text, (f"{dateiname}: Das PDF enthält kaum Text – vermutlich ein Scan. "
                          "Bitte den Text von Hand einfügen.")
        return text, None
    if endung == ".docx":
        text = _text_aus_word(rohdaten)
    elif endung == ".odt":
        text = _text_aus_odt(rohdaten)
    elif endung == ".rtf":
        text = _text_aus_rtf(rohdaten)
    elif endung in (".txt", ".md", ".csv"):
        text = _text_aus_klartext(rohdaten)
    elif endung == ".doc":
        return "", (f"{dateiname}: Das alte Word-Format (.doc) lässt sich nicht lesen. "
                    "Bitte in Word als .docx oder PDF speichern.")
    else:
        return "", f"{dateiname}: Dateityp {endung or '?'} wird nicht gelesen."
    if not text.strip():
        return "", f"{dateiname}: Kein Text gefunden."
    return text, None


# ------------------------------------------------- Improfy-Excel Zelle für Zelle
def _zelle(ws, adresse):
    wert = ws[adresse].value
    if wert is None:
        return ""
    if hasattr(wert, "strftime"):
        return wert.strftime("%d.%m.%Y")
    if isinstance(wert, float) and wert.is_integer():
        wert = int(wert)
    return str(wert).strip()


def _sterne(wert):
    """In der Vorlage stehen Sterne mal als Zahl, mal als '★★★★'."""
    text = str(wert or "").strip()
    if not text:
        return 0
    if text.isdigit():
        return max(1, min(5, int(text)))
    gezaehlt = text.count("★") or text.count("*")
    return max(1, min(5, gezaehlt)) if gezaehlt else 0


def daten_aus_excel(rohdaten):
    """Eine ausgefüllte Improfy-Vorlage zurücklesen. Gibt (daten, gefunden) oder (None, [])."""
    try:
        import openpyxl
        mappe = openpyxl.load_workbook(io.BytesIO(rohdaten), data_only=True)
    except Exception:
        return None, []
    import sys
    hier = os.path.dirname(os.path.abspath(__file__))
    if os.path.join(hier, "cv") not in sys.path:
        sys.path.insert(0, os.path.join(hier, "cv"))
    import fill_cv

    # Das richtige Blatt: das erste, in dem der Namensanker gefüllt ist.
    blatt = None
    for ws in mappe.worksheets:
        if _zelle(ws, fill_cv.KOPF["vorname"]) or _zelle(ws, fill_cv.KOPF["nachname"]):
            blatt = ws
            break
    if blatt is None:
        return None, []

    daten = {"vorname": "", "nachname": "", "geschlecht": "", "angestrebter_job": "",
             "geburtsdatum": "", "mobil": "", "email": "", "adresse": "",
             "kunde_von": _zelle(blatt, fill_cv.KUNDE_VON_ZELLE),
             "fuehrerschein": {"vorhanden": False, "klasse": "", "eu": False},
             "berufserfahrung": [], "bildung": [], "zusatzqualifikationen": [],
             "sprachen": [], "edv_kenntnisse": [], "soft_skills": [],
             "ueber_mich": "", "hobbys": ""}
    for feld, adresse in fill_cv.KOPF.items():
        if feld != "fuehrerschein":
            daten[feld] = _zelle(blatt, adresse)
    fs = _zelle(blatt, fill_cv.KOPF["fuehrerschein"])
    if fs and not re.match(r"^(nein|kein|-|0)$", fs, re.I):
        klasse = re.search(r"\b(?:Klasse\s*)?([A-Z]{1,2}\d?)\b", fs)
        # Die Vorlage füllt B20 nur bei einem Führerschein von außerhalb der EU.
        daten["fuehrerschein"] = {"vorhanden": True,
                                  "klasse": klasse.group(1) if klasse else "",
                                  "eu": not _zelle(blatt, "B20")}

    for zeitraum_z, firma_z, titel_z, taet_z in fill_cv.BERUF_SLOTS:
        zeitraum, firma, titel = (_zelle(blatt, zeitraum_z), _zelle(blatt, firma_z),
                                  _zelle(blatt, titel_z))
        taetigkeiten = [t for t in (_zelle(blatt, z) for z in taet_z) if t]
        if zeitraum or firma or titel:
            daten["berufserfahrung"].append({"zeitraum": zeitraum, "firma": firma,
                                             "jobtitel": titel, "taetigkeiten": taetigkeiten})
    for zeitraum_z, abschluss_z, inst_z, note_z in fill_cv.BILDUNG_SLOTS:
        werte = [_zelle(blatt, z) for z in (zeitraum_z, abschluss_z, inst_z, note_z)]
        if any(werte[:3]):
            daten["bildung"].append(dict(zip(("zeitraum", "abschluss", "institution", "note"), werte)))
    daten["zusatzqualifikationen"] = [z for z in (_zelle(blatt, a) for a in fill_cv.ZUSATZQUAL_ZELLEN) if z]
    for programm_z, sterne_z in fill_cv.EDV_SLOTS:
        programm = _zelle(blatt, programm_z)
        if programm:
            daten["edv_kenntnisse"].append({"programm": programm,
                                            "sterne": _sterne(blatt[sterne_z].value) or 4})
    for zeile in range(fill_cv.SOFT_SKILL_START,
                       fill_cv.SOFT_SKILL_START + fill_cv.SOFT_SKILL_COUNT):
        eigenschaft = _zelle(blatt, f"{fill_cv.SOFT_SKILL_LABEL_COL}{zeile}")
        if eigenschaft:
            daten["soft_skills"].append({"eigenschaft": eigenschaft,
                                         "sterne": _sterne(blatt[f"E{zeile}"].value) or 5})
    for sprache_z, niveau_z in fill_cv.SPRACHE_SLOTS:
        sprache = _zelle(blatt, sprache_z)
        if sprache:
            daten["sprachen"].append({"sprache": sprache.upper(), "niveau": _zelle(blatt, niveau_z)})
    daten["ueber_mich"] = _zelle(blatt, fill_cv.UEBER_MICH_ZELLE)

    gefunden = []
    if daten["vorname"] or daten["nachname"]:
        gefunden.append("Name und Kontakt")
    for schluessel, wort in (("berufserfahrung", "Stationen Berufserfahrung"),
                             ("bildung", "Einträge Bildung"), ("sprachen", "Sprachen"),
                             ("edv_kenntnisse", "EDV-Kenntnisse"), ("soft_skills", "Soft Skills")):
        if daten[schluessel]:
            gefunden.append(f"{len(daten[schluessel])} {wort}")
    if daten["ueber_mich"]:
        gefunden.append("Über mich")
    if daten["fuehrerschein"]["vorhanden"]:
        gefunden.append("Führerschein")
    return (daten, gefunden) if gefunden else (None, [])


# ------------------------------------------------------------------- Zusammenführen
LISTEN = ("berufserfahrung", "bildung", "sprachen", "edv_kenntnisse", "soft_skills",
          "zusatzqualifikationen")


def leere_daten():
    """Alle Felder, die das Formular kennt – auch die, die keine Quelle geliefert hat.

    Sonst fehlt im Ergebnis ein Schlüssel, und die Seite bricht beim Anzeigen ab."""
    return {"kunde_von": "", "vorname": "", "nachname": "", "geschlecht": "",
            "angestrebter_job": "", "geburtsdatum": "", "massnahme_zeitraum": "",
            "mobil": "", "email": "", "adresse": "", "ueber_mich": "", "hobbys": "",
            "fuehrerschein": {"vorhanden": False, "klasse": "", "eu": False},
            "berufserfahrung": [], "bildung": [], "zusatzqualifikationen": [],
            "sprachen": [], "edv_kenntnisse": [], "soft_skills": []}


def _uebernehmen(ziel, quelle):
    """Lücken füllen, nie überschreiben – die erste Quelle ist die verlässlichste."""
    for feld, wert in (quelle or {}).items():
        if feld == "fuehrerschein":
            if wert and wert.get("vorhanden") and not (ziel.get(feld) or {}).get("vorhanden"):
                ziel[feld] = wert
        elif feld in LISTEN:
            if wert and not ziel.get(feld):
                ziel[feld] = wert
        elif wert and not ziel.get(feld):
            ziel[feld] = wert
    return ziel


def lese_unterlagen(dateien, zusatztext=""):
    """Hochgeladene Unterlagen auswerten.

    `dateien` ist eine Liste von (dateiname, rohdaten). Gibt (daten, gefunden, hinweise,
    rohtext) zurück. `rohtext` ist der zusammengesetzte Text, damit er im Textfeld
    stehen bleibt und man ihn nachbessern kann."""
    daten = leere_daten()
    gefunden, hinweise, texte = [], [], []

    # Erst die Excel-Vorlagen: sie sind exakt und gewinnen gegen jede Textauslese.
    for name, rohdaten in dateien:
        if os.path.splitext(name or "")[1].lower() in (".xlsx", ".xlsm"):
            aus_excel, gef = daten_aus_excel(rohdaten)
            if aus_excel:
                daten = _uebernehmen(daten, aus_excel)
                gefunden.append(f"{name}: Improfy-Vorlage gelesen – " + ", ".join(gef))
            else:
                hinweise.append(f"{name}: keine ausgefüllte Improfy-Vorlage.")

    for name, rohdaten in dateien:
        if os.path.splitext(name or "")[1].lower() in (".xlsx", ".xlsm"):
            continue
        text, hinweis = text_aus_datei(name, rohdaten)
        if hinweis:
            hinweise.append(hinweis)
        if text.strip():
            texte.append(text)
            aus_text, gef = cv_lesen.lese_lebenslauf(text)
            if gef:
                daten = _uebernehmen(daten, aus_text)
                gefunden.append(f"{name}: " + ", ".join(gef))
            else:
                hinweise.append(f"{name}: Text gelesen, aber kein bekannter Aufbau erkannt.")

    if (zusatztext or "").strip():
        texte.append(zusatztext)
        aus_text, gef = cv_lesen.lese_lebenslauf(zusatztext)
        if gef:
            daten = _uebernehmen(daten, aus_text)
            gefunden.append("Eingefügter Text: " + ", ".join(gef))

    return (daten if gefunden else None), gefunden, hinweise, "\n\n".join(texte).strip()
